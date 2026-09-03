# -*- coding: utf-8 -*-
import unittest

from pptx import Presentation
from pptx.util import Inches

from app.report.ppt.exceptions import TemplateShapeDuplicateError, TemplateShapeMissingError
from app.report.ppt.shape_resolver import ShapeResolver, get_shape


def make_slide():
    prs = Presentation()
    return prs.slides.add_slide(prs.slide_layouts[6])


def add_named_textbox(shapes, name):
    shape = shapes.add_textbox(Inches(0), Inches(0), Inches(1), Inches(0.3))
    shape.name = name
    return shape


class ShapeResolverTests(unittest.TestCase):
    def test_get_shape_finds_exact_name(self):
        slide = make_slide()
        expected = add_named_textbox(slide.shapes, "ISSUE_TITLE")
        add_named_textbox(slide.shapes, "ISSUE_TITLE_SUFFIX")

        actual = get_shape(slide, "ISSUE_TITLE", slide_key="ISSUE_STANDARD")

        self.assertIs(expected._element, actual._element)

    def test_missing_shape_has_slide_and_shape_context(self):
        slide = make_slide()
        with self.assertRaisesRegex(
            TemplateShapeMissingError,
            "slide=ISSUE_STANDARD; shape=ISSUE_TITLE",
        ):
            get_shape(slide, "ISSUE_TITLE", slide_key="ISSUE_STANDARD")

    def test_duplicate_shape_is_not_silently_selected(self):
        slide = make_slide()
        add_named_textbox(slide.shapes, "ISSUE_TITLE")
        add_named_textbox(slide.shapes, "ISSUE_TITLE")
        with self.assertRaises(TemplateShapeDuplicateError):
            ShapeResolver().get_shape(slide, "ISSUE_TITLE", slide_key="ISSUE_STANDARD")

    def test_finds_shape_nested_in_group(self):
        slide = make_slide()
        group = slide.shapes.add_group_shape()
        expected = add_named_textbox(group.shapes, "IMAGE_MAIN")

        actual = get_shape(slide, "IMAGE_MAIN", slide_key="ISSUE_STANDARD")

        self.assertIs(expected._element, actual._element)


if __name__ == "__main__":
    unittest.main()
