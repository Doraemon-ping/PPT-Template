# -*- coding: utf-8 -*-
import io
import unittest
from pathlib import Path
from unittest.mock import patch

from pptx import Presentation

from app.demo import demo_state

try:
    from fastapi.testclient import TestClient
    from app.main import app
except ModuleNotFoundError:  # Local bundled test runtime may omit web dependencies.
    TestClient = None
    app = None


@unittest.skipUnless(TestClient is not None, "FastAPI/httpx test dependencies are not installed")
class TemplateAPITests(unittest.TestCase):
    def test_formula_preview_renders_existing_parameters(self):
        response = self.client.post('/api/template/formula-preview', json={
            'expression': 'F = [[{f.aPart} × {f.castP}|10]] = {ppt.force.part} kN', 'data': {'f': {'aPart': 820, 'castP': 80}},
        })
        self.assertEqual(200, response.status_code)
        self.assertTrue(response.content.startswith(b'\x89PNG'))
        self.assertEqual('0', response.headers['x-dfm-formula-missing'])

    def test_formula_preview_reports_missing_parameters(self):
        response = self.client.post('/api/template/formula-preview', json={'expression': 'F = {f.notFilled}', 'data': {}})
        self.assertEqual(200, response.status_code)
        self.assertEqual('1', response.headers['x-dfm-formula-missing'])

    def test_formula_preview_rejects_invalid_structure_and_object_values(self):
        for expression in ('', '[[12|]]', '[[[[12|2]]|10]]', '{f}'):
            with self.subTest(expression=expression):
                response = self.client.post('/api/template/formula-preview', json={'expression': expression, 'data': {}})
                self.assertEqual(422, response.status_code)

    @classmethod
    def setUpClass(cls):
        cls.client = TestClient(app)
        cls.payload = demo_state()

    def test_templates_list_includes_registry(self):
        response = self.client.get("/api/templates")
        self.assertEqual(200, response.status_code)
        ids = {item["template_id"] for item in response.json()["templates"]}
        self.assertIn("demo", ids)
        self.assertIn("official", ids)
        self.assertIn("table-demo", ids)

    def test_scan_demo_template_returns_inventory(self):
        response = self.client.post("/api/template/scan", json={"template": "demo"})
        self.assertEqual(200, response.status_code)
        data = response.json()
        self.assertEqual(3, data["slide_count"])
        self.assertEqual(16, data["placeholder_count"])
        self.assertIn("f.partNo", data["placeholder_paths"])

    def test_scan_unknown_template_returns_404(self):
        response = self.client.post("/api/template/scan", json={"template": "nope"})
        self.assertEqual(404, response.status_code)

    def test_generate_demo_deck_returns_pptx_with_headers(self):
        payload = {
            "template": "demo",
            "output_mode": "deck",
            "missing": "keep",
            "slides": [
                {"source": 1},
                {"source": 2, "images": {"PART_IMAGE": "i.logoImg[0]"}},
                {"source": 3, "repeat": "t.issues"},
            ],
            "data": self.payload,
        }
        response = self.client.post("/api/template/generate", json=payload)
        self.assertEqual(200, response.status_code)
        self.assertEqual("openxml-template-v1", response.headers["x-dfm-engine"])
        self.assertEqual("4", response.headers["x-dfm-slide-count"])
        self.assertEqual("21", response.headers["x-dfm-text-replaced"])
        self.assertEqual("1", response.headers["x-dfm-images-bound"])
        self.assertIn("DFM_TEMPLATE_", response.headers["content-disposition"])
        self.assertEqual(4, len(Presentation(io.BytesIO(response.content)).slides))

    def test_generate_rejects_bad_output_mode(self):
        response = self.client.post(
            "/api/template/generate",
            json={"template": "demo", "output_mode": "bogus", "slides": [{"source": 1}], "data": {}},
        )
        self.assertEqual(422, response.status_code)

    def test_generate_missing_product_images_returns_pptx_and_warning_count(self):
        response = self.client.post('/api/template/generate', json={
            'template': 'demo', 'slides': [{'source': 2, 'bindings': {
                'front': {'type': 'image', 'source': 'i.productRunnerFrontImg[0]', 'shape': 'PART_IMAGE'},
                'back': {'type': 'image_region', 'source': 'i.productRunnerBackImg[0]', 'shape': 'PART_IMAGE'},
            }}], 'data': {'f': {}, 'i': {}},
        })
        self.assertEqual(200, response.status_code)
        self.assertEqual('2', response.headers['x-dfm-images-missing'])
        self.assertEqual(1, len(Presentation(io.BytesIO(response.content)).slides))

    def test_live_preview_applies_data_and_does_not_modify_template(self):
        from app.main import _registry
        original_path = _registry().resolve("demo").path
        original = original_path.read_bytes()
        rendered_paths = []

        def fake_render(pptx_path, slide_index, cache_root):
            deck = Presentation(pptx_path)
            text = "\n".join(shape.text for shape in deck.slides[slide_index - 1].shapes if shape.has_text_frame)
            self.assertIn(self.payload["f"]["partNo"], text)
            target = Path(cache_root) / "preview.png"
            target.parent.mkdir(parents=True)
            target.write_bytes(b"preview-png")
            rendered_paths.append(Path(pptx_path))
            return target

        with patch("app.report.ppt.slide_preview.render_slide_preview", side_effect=fake_render):
            response = self.client.post("/api/template/live-preview", json={
                "template": "demo", "slide": {"source": 1, "condition": "f.missing", "bindings": {
                    "incomplete": {"type": "image", "shape": "Unfilled", "source": "i.unfilled[0]"},
                }}, "data": self.payload,
            })
        self.assertEqual(200, response.status_code, response.text)
        self.assertEqual("image/png", response.headers["content-type"])
        self.assertEqual("no-store", response.headers["cache-control"])
        self.assertEqual("1", response.headers["x-dfm-preview-skipped"])
        self.assertEqual(b"preview-png", response.content)
        self.assertFalse(rendered_paths[0].exists())
        self.assertEqual(original, original_path.read_bytes())

    def test_live_preview_repeated_page_uses_first_record(self):
        def fake_render(pptx_path, slide_index, cache_root):
            deck = Presentation(pptx_path)
            text = "\n".join(shape.text for shape in deck.slides[slide_index - 1].shapes if shape.has_text_frame)
            self.assertIn("PREVIEW_FIRST", text)
            self.assertNotIn("PREVIEW_SECOND", text)
            target = Path(cache_root) / "preview.png"
            target.parent.mkdir(parents=True)
            target.write_bytes(b"preview-png")
            return target

        data = {"t": {"issues": [{"desc": "PREVIEW_FIRST"}, {"desc": "PREVIEW_SECOND"}]}}
        with patch("app.report.ppt.slide_preview.render_slide_preview", side_effect=fake_render):
            response = self.client.post("/api/template/live-preview", json={
                "template": "demo", "slide": {"source": 3, "repeat": "t.issues"}, "data": data,
            })
        self.assertEqual(200, response.status_code, response.text)

    def test_live_preview_reuses_process_cache_for_identical_request(self):
        calls = []

        def fake_render(pptx_path, slide_index, cache_root):
            calls.append((Path(pptx_path), slide_index))
            target = Path(cache_root) / "preview.png"
            target.parent.mkdir(parents=True, exist_ok=True)
            target.write_bytes(b"cached-preview-png")
            return target

        payload = {
            "template": "demo",
            "slide": {"source": 1},
            "data": {"f": {"cache_marker": "unique-live-preview-cache-test"}},
        }
        with patch("app.report.ppt.slide_preview.render_slide_preview", side_effect=fake_render):
            first = self.client.post("/api/template/live-preview", json=payload)
            second = self.client.post("/api/template/live-preview", json=payload)
        self.assertEqual(200, first.status_code, first.text)
        self.assertEqual(200, second.status_code, second.text)
        self.assertEqual(1, len(calls))
        self.assertEqual("miss", first.headers["x-dfm-preview-cache"])
        self.assertEqual("hit", second.headers["x-dfm-preview-cache"])
        self.assertEqual(b"cached-preview-png", second.content)

    def test_live_preview_unknown_template_returns_404(self):
        response = self.client.post("/api/template/live-preview", json={"template": "nope", "slide": {"source": 1}})
        self.assertEqual(404, response.status_code)

    def test_partial_binding_preserves_label_in_generated_deck(self):
        inspected = self.client.post('/api/template/inspect', json={'template': 'demo'}).json()
        shape = next(s for s in inspected['slides'][0]['shapes'] if s['shape_name'] == 'COVER_CUSTOMER')
        original = shape['full_text']
        token = '{f.custName}'
        start = original.index(token)
        binding = {'type': 'text_replace', 'source': 'f', 'shape': shape['shape_name'], 'options': {
            'shape_id': shape['shape_id'], 'original_text': original,
            'replacements': [{'start': start, 'end': start+len(token), 'text': token, 'source': 'f.custName'}],
        }}
        response = self.client.post('/api/template/generate', json={
            'template': 'demo', 'slides': [{'source': 1, 'bindings': {'customer': binding}}],
            'data': {'f': {'custName': '局部测试客户'}},
        })
        self.assertEqual(200, response.status_code, response.text if response.status_code != 200 else '')
        deck = Presentation(io.BytesIO(response.content))
        actual = next(s.text for s in deck.slides[0].shapes if s.name == 'COVER_CUSTOMER')
        self.assertEqual(original.replace(token, '局部测试客户'), actual)

    def test_legacy_ppt_route_remains_unchanged(self):
        response = self.client.post("/api/ppt", json=self.payload)
        self.assertEqual(200, response.status_code)
        self.assertEqual(47, len(Presentation(io.BytesIO(response.content)).slides))


if __name__ == "__main__":
    unittest.main()
