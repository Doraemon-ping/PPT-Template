# -*- coding: utf-8 -*-
import io
import unittest
from pathlib import Path

from pptx import Presentation

from app.demo import demo_state
from app.report.ppt.deck import DeckDefinition, DeckPlanner, DeckSlide
from app.report.ppt.template_engine import TemplateEngine

ROOT = Path(__file__).resolve().parents[1]
DEMO_TEMPLATE = ROOT / "templates" / "DFM_Template_Placeholder_Demo.pptx"
TABLE_TEMPLATE = ROOT / "templates" / "DFM_Template_Table_Demo.pptx"


def demo_deck(*, output_mode="deck", missing="keep", condition=None, include_images=True):
    slides = [DeckSlide(source=1)]
    if include_images:
        slides.append(DeckSlide(source=2, images={"PART_IMAGE": "i.logoImg[0]"}))
    else:
        slides.append(DeckSlide(source=2))
    slides.append(DeckSlide(source=3, repeat="t.issues", condition=condition))
    return DeckDefinition(template="demo", output_mode=output_mode, missing=missing, slides=slides)


class DeckPlannerTests(unittest.TestCase):
    def test_expand_repeat_items_and_condition(self):
        data = demo_state()
        data["f"]["leakReq"] = "气压测试"
        plan = DeckPlanner().expand(
            DeckDefinition(template="demo", slides=[
                DeckSlide(source=1),
                DeckSlide(source=3, repeat="t.issues", condition="f.leakReq"),
            ]),
            data,
        )
        self.assertEqual(2, len(plan.ops))
        self.assertEqual(2, len(plan.ops[1].items))  # 2 条示例 issue

        empty = dict(data)
        empty["f"] = dict(data["f"])
        empty["f"]["leakReq"] = ""
        plan2 = DeckPlanner().expand(
            DeckDefinition(template="demo", slides=[
                DeckSlide(source=3, repeat="t.issues", condition="f.leakReq"),
            ]),
            empty,
        )
        self.assertEqual(0, len(plan2.ops))  # 条件不满足 -> 整页跳过


class TemplateEngineTests(unittest.TestCase):
    def test_deck_mode_generates_repeat_slides_with_filled_text_and_image(self):
        result = TemplateEngine(DEMO_TEMPLATE).generate(demo_deck(), demo_state())
        self.assertEqual(4, result.slide_count)
        self.assertEqual(21, result.stats["text_replaced"])
        self.assertEqual([], result.stats["text_missing"])
        self.assertEqual(1, result.stats["images_bound"])
        prs = Presentation(result.bytes_io)
        self.assertEqual(4, len(prs.slides))
        self.assertIn("TP-HPDC-2026-0087", prs.slides[0].shapes[0].text)
        issue_texts = [sh.text for sh in prs.slides[3].shapes if getattr(sh, "has_text_frame", False)]
        self.assertTrue(any("反拔模区域需滑块抽芯" in text for text in issue_texts))

    def test_in_place_mode_keeps_all_slides_and_fills_cover(self):
        deck = DeckDefinition(template="demo", output_mode="in_place", slides=[DeckSlide(source=1)])
        result = TemplateEngine(DEMO_TEMPLATE).generate(deck, demo_state())
        self.assertEqual(3, result.slide_count)
        prs = Presentation(result.bytes_io)
        self.assertEqual(3, len(prs.slides))
        self.assertIn("TP-HPDC-2026-0087", prs.slides[0].shapes[0].text)
        # 未选中的第 2 页保持模板原样（仍含占位符）
        slide2_texts = [sh.text for sh in prs.slides[1].shapes if getattr(sh, "has_text_frame", False)]
        self.assertTrue(any("{f.wCast}" in text for text in slide2_texts))

    def test_missing_policy_clear_removes_unbound_placeholders(self):
        deck = DeckDefinition(template="demo", output_mode="deck", missing="clear", slides=[DeckSlide(source=1)])
        result = TemplateEngine(DEMO_TEMPLATE).generate(deck, {"f": {}, "t": {}, "i": {}})
        prs = Presentation(result.bytes_io)
        texts = [sh.text for sh in prs.slides[0].shapes if getattr(sh, "has_text_frame", False)]
        self.assertTrue(any("{" not in text for text in texts if text.strip()))

    def test_missing_source_slide_raises(self):
        deck = DeckDefinition(template="demo", slides=[DeckSlide(source=99)])
        with self.assertRaises(Exception):
            TemplateEngine(DEMO_TEMPLATE).generate(deck, demo_state())

    def test_unbound_package_parts_preserved_byte_identical(self):
        import zipfile

        result = TemplateEngine(DEMO_TEMPLATE).generate(demo_deck(), demo_state())
        with zipfile.ZipFile(io.BytesIO(result.buffer)) as out_zip, \
                zipfile.ZipFile(DEMO_TEMPLATE) as in_zip:
            for name in ("ppt/theme/theme1.xml", "ppt/slideMasters/slideMaster1.xml"):
                self.assertEqual(in_zip.read(name), out_zip.read(name), name)

    def test_multi_template_deck_combines_slides_in_one_output(self):
        """跨模板拼页：demo 第1页 + table-demo 第1页 + demo 第3页(重复)。"""
        deck = DeckDefinition(template="demo", output_mode="deck", slides=[
            DeckSlide(source=1),
            DeckSlide(source=1, template="table-demo", bindings={
                "c1": {"type": "table_cell", "source": "f.material",
                       "shape": "PART_TABLE", "options": {"row": 4, "column": 1}},
            }),
            DeckSlide(source=3, repeat="t.issues"),
        ])
        result = TemplateEngine(DEMO_TEMPLATE).generate(
            deck, demo_state(),
            template_map={"demo": DEMO_TEMPLATE, "table-demo": TABLE_TEMPLATE},
        )
        self.assertEqual(4, result.slide_count)
        prs = Presentation(result.bytes_io)
        self.assertEqual(4, len(prs.slides))
        cover_texts = [sh.text for sh in prs.slides[0].shapes if getattr(sh, "has_text_frame", False)]
        self.assertTrue(any("TP-HPDC-2026-0087" in t for t in cover_texts))
        # 第 2 页是导入的表格模板页：占位符 + 显式单元格绑定均生效
        imported = prs.slides[1]
        tables = [sh for sh in imported.shapes if getattr(sh, "has_table", False)]
        self.assertEqual(1, len(tables))
        self.assertEqual("AlSi10MnMg", tables[0].table.cell(4, 1).text.strip())
        issue_texts = [sh.text for sh in prs.slides[3].shapes if getattr(sh, "has_text_frame", False)]
        self.assertTrue(any("反拔模" in t for t in issue_texts))


if __name__ == "__main__":
    unittest.main()
