# -*- coding: utf-8 -*-
import io
import os
import unittest
from unittest.mock import patch

from pptx import Presentation

from app.demo import demo_state

try:
    from fastapi.testclient import TestClient
    from app.main import app
    from app.dfm import adapt_legacy_report
    from app.report.ppt.engine import GENERATOR_VERSION
    from app.report.ppt.slide_planner import plan_slides
except ModuleNotFoundError:  # Local bundled test runtime may omit web dependencies.
    TestClient = None
    app = None
    GENERATOR_VERSION = None
    adapt_legacy_report = None
    plan_slides = None


@unittest.skipUnless(TestClient is not None, "FastAPI/httpx test dependencies are not installed")
class PPTPreviewAPITests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.client = TestClient(app)
        cls.payload = demo_state()

    def test_preview_route_is_disabled_by_default(self):
        environment = dict(os.environ)
        environment.pop("DFM_PPT_V2_ENABLED", None)
        with patch.dict(os.environ, environment, clear=True):
            response = self.client.post("/api/ppt/preview-v2", json=self.payload)

        self.assertEqual(503, response.status_code)
        self.assertEqual("PPT V2 预览功能未启用", response.json()["detail"])

    def test_enabled_preview_returns_valid_ppt_and_diagnostic_headers(self):
        with patch.dict(os.environ, {"DFM_PPT_V2_ENABLED": "true"}, clear=False):
            response = self.client.post("/api/ppt/preview-v2", json=self.payload)

        self.assertEqual(200, response.status_code)
        self.assertEqual("preview-v2", response.headers["x-dfm-engine"])
        self.assertEqual("1", response.headers["x-dfm-template-version"])
        self.assertEqual(GENERATOR_VERSION, response.headers["x-dfm-generator-version"])
        expected_plans = plan_slides(adapt_legacy_report(**self.payload))
        expected_slide_count = str(len(expected_plans))
        self.assertEqual(expected_slide_count, response.headers["x-dfm-slide-count"])
        self.assertEqual("2", response.headers["x-dfm-validation-warnings"])
        self.assertIn("DFM_PREVIEW_V2_", response.headers["content-disposition"])
        self.assertEqual(
            len(expected_plans),
            len(Presentation(io.BytesIO(response.content)).slides),
        )

    def test_enabled_preview_reports_engine_stage_on_failure(self):
        with patch.dict(
            os.environ,
            {
                "DFM_PPT_V2_ENABLED": "true",
                "DFM_PPT_V2_TEMPLATE": "missing-template.pptx",
            },
            clear=False,
        ):
            response = self.client.post("/api/ppt/preview-v2", json=self.payload)

        self.assertEqual(500, response.status_code)
        self.assertIn("stage=load", response.json()["detail"])

    def test_legacy_ppt_route_remains_unchanged(self):
        response = self.client.post("/api/ppt", json=self.payload)

        self.assertEqual(200, response.status_code)
        self.assertNotIn("x-dfm-engine", response.headers)
        self.assertEqual(47, len(Presentation(io.BytesIO(response.content)).slides))


if __name__ == "__main__":
    unittest.main()
