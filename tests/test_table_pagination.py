import io
import posixpath
import unittest
from copy import deepcopy
from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace
from unittest.mock import patch

from lxml import etree
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN

from app.report.ppt.deck import DeckDefinition, DeckSlide
from app.report.ppt.template_engine import TemplateEngine, TemplateEngineError
from app.report.ppt.openxml.package_editor import OoxmlPackage
from app.report.ppt.openxml.shape_binding import NS


def fixture():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    shape = slide.shapes.add_table(5, 2, Inches(.5), Inches(1), Inches(8), Inches(2.5))
    shape.name = 'DATA'
    shape.table.cell(0, 0).text = '编号'
    shape.table.cell(0, 1).text = '描述'
    for i, height in enumerate([.5, .5, .15, .85, .5]):
        shape.table.rows[i].height = Inches(height)
    for cell in shape.table.rows[1].cells:
        cell.text = '参考'
        cell.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
        cell.text_frame.paragraphs[0].runs[0].font.size = Pt(14)
    shape.table.cell(2, 0).text = '旧数据'
    shape.table.cell(2, 0).text_frame.paragraphs[0].alignment = PP_ALIGN.LEFT
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def spec(**options):
    return DeckSlide(source=1, bindings={'rows': {'type': 'table_rows', 'shape': 'DATA', 'source': 't.rows',
                        'options': {'columns_keys': ['id', 'text'], 'start_row': 1, **options}}})


def data(n, text='测试'):
    return {'t': {'rows': [{'id': str(i), 'text': text} for i in range(n)]}}


class TablePaginationTests(unittest.TestCase):
    def generate(self, n, **options):
        source = fixture()
        deck = DeckDefinition(template='test', slides=[spec(**options)])
        before = deck.model_dump()
        result = TemplateEngine(source).generate(deck, data(n))
        self.assertEqual(before, deck.model_dump())
        return Presentation(result.bytes_io), result

    def test_uniform_rows_and_all_records_exactly_once(self):
        prs, result = self.generate(11)
        self.assertEqual(3, len(prs.slides))
        ids, heights = [], set()
        for slide in prs.slides:
            table = slide.shapes[0].table
            self.assertEqual('编号', table.cell(0, 0).text)
            self.assertLessEqual(sum(r.height for r in table.rows), Inches(2.5))
            for i in range(1, len(table.rows)):
                ids.append(table.cell(i, 0).text)
                heights.add(table.rows[i].height)
                self.assertEqual(PP_ALIGN.CENTER, table.cell(i, 0).text_frame.paragraphs[0].alignment)
        self.assertEqual([str(i) for i in range(11)], ids)
        self.assertEqual({Inches(.5)}, heights)

    def test_empty_and_exact_capacity(self):
        for n, rows in [(0, 1), (4, 5)]:
            prs, _ = self.generate(n)
            self.assertEqual(1, len(prs.slides))
            self.assertEqual(rows, len(prs.slides[0].shapes[0].table.rows))

    def test_explicit_page_limit(self):
        prs, _ = self.generate(5, rows_per_page=2)
        self.assertEqual(3, len(prs.slides))
        self.assertEqual([3, 3, 2], [len(s.shapes[0].table.rows) for s in prs.slides])
        self.assertEqual({Inches(1)}, {r.height for s in prs.slides for r in list(s.shapes[0].table.rows)[1:]})

    def test_explicit_count_can_exceed_template_capacity(self):
        prs, _ = self.generate(13, rows_per_page=5)
        self.assertEqual([6, 6, 4], [len(s.shapes[0].table.rows) for s in prs.slides])
        self.assertEqual({Inches(2) // 5}, {r.height for s in prs.slides for r in list(s.shapes[0].table.rows)[1:]})
        ids = [s.shapes[0].table.cell(i, 0).text for s in prs.slides for i in range(1, len(s.shapes[0].table.rows))]
        self.assertEqual([str(i) for i in range(13)], ids)
        self.assertEqual(Pt(14), prs.slides[0].shapes[0].table.cell(1, 0).text_frame.paragraphs[0].runs[0].font.size)

    def test_explicit_count_too_dense_is_not_silently_ignored(self):
        with self.assertRaisesRegex(TemplateEngineError, '每页 20 行放不下.*最多可放'):
            self.generate(25, rows_per_page=20)

    def test_invalid_page_counts(self):
        for limit in (-1, 1.5, '2.5', 'abc', True, [], {}):
            with self.subTest(limit=limit), self.assertRaisesRegex(TemplateEngineError, '每页条数'):
                self.generate(5, rows_per_page=limit)

    def test_preview_uses_explicit_count(self):
        result = TemplateEngine(fixture()).generate(
            DeckDefinition(template='test', output_mode='in_place', slides=[spec(rows_per_page=5)]),
            data(13), preview_page=2)
        self.assertEqual(3, result.stats['preview_page_count'])
        table = Presentation(result.bytes_io).slides[0].shapes[0].table
        self.assertEqual(6, len(table.rows))
        self.assertEqual('5', table.cell(1, 0).text)

    def test_style_reference_selection(self):
        prs, _ = self.generate(6, style_row=3)
        self.assertEqual(3, len(prs.slides))
        self.assertEqual(Inches(.85), prs.slides[0].shapes[0].table.rows[1].height)

    def test_multiline_increases_uniform_height_and_page_count(self):
        result = TemplateEngine(fixture()).generate(DeckDefinition(template='test', slides=[spec()]), data(4, '长文字\n第二行\n第三行\n第四行'))
        prs = Presentation(result.bytes_io)
        self.assertGreater(len(prs.slides), 1)
        heights = {r.height for s in prs.slides for r in list(s.shapes[0].table.rows)[1:]}
        self.assertEqual(1, len(heights))
        self.assertGreater(next(iter(heights)), Inches(.5))

    def test_oversized_record_fails_without_truncation(self):
        with self.assertRaisesRegex(TemplateEngineError, '单条表格内容'):
            TemplateEngine(fixture()).generate(DeckDefinition(template='test', slides=[spec()]), data(1, '长文字' * 2000))

    def test_repeat_and_pagination_keep_item_scope(self):
        page = spec()
        page.repeat = 't.groups'
        page.bindings['rows'].source = 'rows'
        result = TemplateEngine(fixture()).generate(DeckDefinition(template='test', slides=[page]),
            {'t': {'groups': [{'rows': data(5)['t']['rows']}, {'rows': [{'id': 'NEXT', 'text': '组2'}]}]}})
        prs = Presentation(result.bytes_io)
        self.assertEqual(3, len(prs.slides))
        self.assertEqual('NEXT', prs.slides[2].shapes[0].table.cell(1, 0).text)

    def test_cross_template_continuation(self):
        page = spec()
        page.template = 'other'
        result = TemplateEngine(fixture()).generate(DeckDefinition(template='base', slides=[page]), data(5), template_map={'other': fixture()})
        self.assertEqual(2, result.slide_count)

    def test_in_place_rejects_extra_pages_but_preview_can_select(self):
        source = fixture()
        deck = DeckDefinition(template='test', output_mode='in_place', slides=[spec()])
        with self.assertRaisesRegex(TemplateEngineError, '续页'):
            TemplateEngine(source).generate(deck, data(5))
        result = TemplateEngine(source).generate(deck, data(5), preview_page=2)
        self.assertEqual(2, result.stats['preview_page_count'])
        self.assertEqual('4', Presentation(result.bytes_io).slides[0].shapes[0].table.cell(1, 0).text)

    def test_invalid_reference(self):
        for index in (0, 12):
            with self.assertRaisesRegex(TemplateEngineError, '参考行'):
                self.generate(2, style_row=index)

    def test_continuations_clone_owned_tags_in_same_and_cross_template(self):
        package = OoxmlPackage(fixture())
        rel_ns = 'http://schemas.openxmlformats.org/package/2006/relationships'
        rel_part = 'ppt/slides/_rels/slide1.xml.rels'
        root = etree.fromstring(package.read(rel_part))
        rel = etree.SubElement(root, '{'+rel_ns+'}Relationship', Id='rId99',
            Type='http://schemas.openxmlformats.org/officeDocument/2006/relationships/tags', Target='../tags/tag1.xml')
        package.write(rel_part, etree.tostring(root))
        package.write('ppt/tags/tag1.xml', b'<p:tagLst xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"/>')
        package.ensure_content_type('ppt/tags/tag1.xml', 'application/vnd.openxmlformats-officedocument.presentationml.tags+xml')
        for cross in (False, True):
            page = spec()
            page.template = 'other' if cross else None
            result = TemplateEngine(package.save()).generate(DeckDefinition(template='test', slides=[page]), data(9), template_map={'other': package.save()})
            out = OoxmlPackage(result.buffer)
            targets = []
            for slide_part in out.slide_parts().values():
                relations = etree.fromstring(out.read(out.rels_part_for(slide_part)))
                for item in relations:
                    if item.get('Type').endswith('/tags'):
                        target = posixpath.normpath(posixpath.join(posixpath.dirname(slide_part), item.get('Target')))
                        self.assertTrue(out.has_part(target))
                        targets.append(target)
                    self.assertFalse(item.get('Type').endswith('/notesSlide'))
            self.assertEqual(3, len(targets))
            self.assertEqual(3, len(set(targets)))
            master_ids = []
            presentation = etree.fromstring(out.read('ppt/presentation.xml'))
            master_ids.extend(presentation.xpath('.//p:sldMasterId/@id', namespaces=NS))
            for part in out.part_names():
                if part.startswith('ppt/slideMasters/') and part.endswith('.xml'):
                    master_ids.extend(etree.fromstring(out.read(part)).xpath('.//p:sldLayoutId/@id', namespaces=NS))
            self.assertEqual(len(master_ids), len(set(master_ids)))

    def test_merged_header_is_preserved(self):
        package = OoxmlPackage(fixture())
        root = etree.fromstring(package.read('ppt/slides/slide1.xml'))
        header = root.find('.//a:tr', NS)
        header.find('a:tc', NS).set('gridSpan', '2')
        header.findall('a:tc', NS)[1].set('hMerge', '1')
        before = etree.tostring(header)
        package.write('ppt/slides/slide1.xml', etree.tostring(root))
        result = TemplateEngine(package.save()).generate(DeckDefinition(template='test', slides=[spec()]), data(5))
        out = OoxmlPackage(result.buffer)
        for part in out.slide_parts().values():
            self.assertEqual(before, etree.tostring(etree.fromstring(out.read(part)).find('.//a:tr', NS)))

    def test_preview_api_returns_continuation_count(self):
        from fastapi.testclient import TestClient
        from app.main import app
        with TemporaryDirectory() as tmp:
            template = Path(tmp) / 'fixture.pptx'
            template.write_bytes(fixture())
            def render(path, slide_index, cache):
                prs = Presentation(path)
                self.assertEqual('4', prs.slides[0].shapes[0].table.cell(1, 0).text)
                target = Path(tmp) / 'preview.png'
                target.write_bytes(b'png')
                return target
            registry = SimpleNamespace(resolve=lambda _: SimpleNamespace(path=template))
            with patch('app.main._registry', return_value=registry), patch('app.report.ppt.slide_preview.render_slide_preview', side_effect=render):
                response = TestClient(app).post('/api/template/live-preview', json={'template': 'test', 'slide': spec().model_dump(), 'data': data(5), 'page': 2})
            self.assertEqual(200, response.status_code, response.text)
            self.assertEqual('2', response.headers['x-dfm-preview-pages'])
            self.assertEqual('2', response.headers['x-dfm-preview-page'])


if __name__ == '__main__':
    unittest.main()
