"""Real independent processes, separate data roots, no in-process gateway."""
import io
import json
import os
import socket
import subprocess
import sys
import tempfile
import time
import unittest
from pathlib import Path

import httpx
from pptx import Presentation
from pptx.util import Inches


class IndependentHTTPTests(unittest.TestCase):
    def test_three_services_and_new_provider(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            sockets = [socket.socket() for _ in range(4)]
            for sock in sockets:
                sock.bind(('127.0.0.1', 0))
            ports = [sock.getsockname()[1] for sock in sockets]
            for sock in sockets:
                sock.close()
            urls = [f'http://127.0.0.1:{port}' for port in ports]
            config = root/'connections.json'
            config.write_text(json.dumps({'connections': [
                {'id': name, 'base_url': url} for name, url in zip(('hpdc','machining','example'), (urls[0],urls[1],urls[3]))]}))
            processes, logs = [], []
            try:
                for i, module in enumerate(('app.services.hpdc','app.services.machining','app.services.workbench','tools.example_ppt_provider')):
                    env = {**os.environ, 'DFM_APP_ROOT': str(root/str(i)), 'PPT_CONNECTIONS_FILE': str(config), 'PPT_WORKBENCH_URL': urls[2], 'PYTHONIOENCODING': 'utf-8'}
                    env.pop('PPT_PROVIDER_TOKEN', None)
                    log = (root/f'{i}.log').open('w', encoding='utf-8'); logs.append(log)
                    processes.append(subprocess.Popen([sys.executable, '-m', 'uvicorn', module+':app', '--host','127.0.0.1','--port',str(ports[i])],
                        cwd=Path(__file__).resolve().parents[1], env=env, stdout=log, stderr=subprocess.STDOUT,
                        creationflags=getattr(subprocess, 'CREATE_NO_WINDOW', 0)))
                with httpx.Client(timeout=30, trust_env=False) as client:
                    for url in urls:
                        deadline = time.monotonic()+25
                        while True:
                            try:
                                if client.get(url+'/health').status_code == 200:
                                    break
                            except httpx.TransportError:
                                pass
                            if time.monotonic() > deadline:
                                self.fail('Service failed to start: '+url)
                            time.sleep(.1)
                    sources = client.get(urls[2]+'/api/ppt/sources').json()
                    self.assertEqual([], sources['errors'])
                    self.assertEqual({'dfm','machining-dfm','inspection'}, {s['id'] for s in sources['sources']})
                    hpdc = client.post(urls[0]+'/api/form-apps/dfm/projects',json={'name':'HTTP test','data':{'f':{'partNo':'HPDC-HTTP'},'t':{},'i':{}}}).json()
                    machine = client.get(urls[1]+'/api/machining-dfm/bootstrap').json()['project']
                    for source, project_id, field, expected in (
                        ('dfm',hpdc['id'],'f.partNo','HPDC-HTTP'),
                        ('machining-dfm',machine['id'],None,machine['state']['G']['part']),
                        ('inspection','sample','f.part','INSPECT-001')):
                        snap = client.get(f'{urls[2]}/api/ppt/sources/{source}/projects/{project_id}/snapshot')
                        self.assertEqual(200,snap.status_code,snap.text[:300]); snap=snap.json()
                        if field is None:
                            field=next(f['path'] for f in snap['catalog']['fields'] if f.get('source_path')=='G.part')
                        deck=Presentation(); slide=deck.slides.add_slide(deck.slide_layouts[6])
                        shape=slide.shapes.add_textbox(Inches(1),Inches(1),Inches(8),Inches(1));shape.text='{'+field+'}'
                        content=io.BytesIO();deck.save(content)
                        upload=client.post(urls[2]+'/api/templates/upload',params={'app_id':source,'template_id':'http-test'},files={'file':('test.pptx',content.getvalue())})
                        self.assertEqual(200,upload.status_code,upload.text)
                        result=client.post(urls[2]+'/api/template/generate',params={'app_id':source},json={
                            'template':'http-test','slides':[{'source':1}],'data':snap['data']})
                        self.assertEqual(200,result.status_code,result.text[:300] if result.status_code!=200 else '')
                        self.assertEqual(expected,Presentation(io.BytesIO(result.content)).slides[0].shapes[0].text)
                        self.assertTrue((root/'2/data/ppt_workbench/sources'/source/'templates/http-test/master.pptx').is_file())
                    link=client.get(urls[0]+'/api/integration/workbench-link',params={'project_id':hpdc['id']}).json()['url']
                    self.assertTrue(link.startswith(urls[2]+'/template-editor'))
                    self.assertEqual(200,client.get(link).status_code)
                    self.assertEqual(200,client.get(urls[0]+'/api/schemes').status_code)
                    self.assertFalse((root/'2/data/form_platform').exists())
                    self.assertFalse((root/'2/data/machining_dfm').exists())
                    self.assertFalse((root/'0/data/ppt_workbench').exists())
                    processes[1].terminate();processes[1].wait(timeout=10)
                    listing=client.get(urls[2]+'/api/ppt/sources').json()
                    self.assertEqual({'dfm','inspection'}, {s['id'] for s in listing['sources']})
                    self.assertEqual('machining',listing['errors'][0]['connection'])
            finally:
                for process in processes:
                    if process.poll() is None:
                        process.terminate()
                    try:
                        process.wait(timeout=10)
                    except subprocess.TimeoutExpired:
                        process.kill();process.wait(timeout=5)
                for log in logs:
                    log.close()
