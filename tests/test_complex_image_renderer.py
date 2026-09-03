# -*- coding: utf-8 -*-
import io
import unittest

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches

from app.report.ppt.exceptions import RendererError
from app.report.ppt.renderers import (
    ComplexImagePayload,
    ComplexImageRenderer,
    RenderContext,
    SnapshotBackend,
)


def png_bytes(width=400, height=200):
    stream = io.BytesIO()
    Image.new("RGB", (width, height), "red").save(stream, format="PNG")
    return stream.getvalue()


def make_frame():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    frame = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(4), Inches(3)
    )
    frame.name = "IMAGE_MAIN"
    return slide, frame


class FakeSnapshotBackend:
    def render_html(self, html, *, width_px, height_px):
        return png_bytes(width_px, height_px)


class ComplexImageRendererTests(unittest.TestCase):
    def test_accepts_existing_raw_image_without_recomposition(self):
        slide, frame = make_frame()
        result = ComplexImageRenderer().render(
            frame,
            png_bytes(),
            RenderContext(slide_key="ISSUE_STANDARD", options={"mode": "contain"}),
        )
        self.assertEqual((400, 200), (result.image_width_px, result.image_height_px))
        self.assertEqual("IMAGE_MAIN", slide.shapes[0].name)
        self.assertAlmostEqual(2.0, result.width / result.height, places=3)

    def test_accepts_explicit_precomposited_payload(self):
        _, frame = make_frame()
        result = ComplexImageRenderer().render(
            frame,
            ComplexImagePayload(png_bytes(100, 200), source="legacy-cad"),
        )
        self.assertEqual((100, 200), (result.image_width_px, result.image_height_px))

    def test_rejects_uncomposited_annotation_request(self):
        _, frame = make_frame()
        with self.assertRaisesRegex(RendererError, "annotations_baked_in must be true"):
            ComplexImageRenderer().render(
                frame,
                {"image": png_bytes(), "annotations_baked_in": False},
                RenderContext(slide_key="ISSUE_STANDARD"),
            )

    def test_snapshot_backend_is_a_runtime_checkable_future_boundary(self):
        self.assertIsInstance(FakeSnapshotBackend(), SnapshotBackend)


if __name__ == "__main__":
    unittest.main()
