# -*- coding: utf-8 -*-
import io
import unittest
from pathlib import Path

from pptx import Presentation

from app.demo import demo_state

try:
    from fastapi.testclient import TestClient
    from app.main import app
except ModuleNotFoundError:  # Local bundled test runtime may omit web dependencies.
    TestClient = None
    app = None

ROOT = Path(__file__).resolve().parents[1]
DEMO = ROOT / "templates" / "DFM_Template_Placeholder_Demo.pptx"
TABLE = ROOT / "templates" / "DFM_Template_Table_Demo.pptx"


@unittest.skipUnless(TestClient is not None, "FastAPI/httpx test dependencies are not installed")
class TemplateImportBindAPITests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.client = TestClient(app)
        cls.payload = demo_state()

    def test_upload_scan_inspect_bind_generate_delete_flow(self):
        payload_id = "apitest-demo"
        # 1) 上传
        with open(DEMO, "rb") as fp:
            response = self.client.post(
                f"/api/templates/upload?template_id={payload_id}&version=9",
                files={"file": ("demo.pptx", fp,
                                "application/vnd.openxmlformats-officedocument.presentationml.presentation")},
            )
        self.assertEqual(200, response.status_code, response.text)
        self.assertEqual(payload_id, response.json()["template"]["template_id"])
        self.assertEqual(3, response.json()["slide_count"])

        # 2) 列表包含新模板
        listed = {t["template_id"] for t in self.client.get("/api/templates").json()["templates"]}
        self.assertIn(payload_id, listed)

        # 3) 扫描
        scan = self.client.post("/api/template/scan", json={"template": payload_id})
        self.assertEqual(200, scan.status_code)
        self.assertEqual(16, scan.json()["placeholder_count"])

        # 4) 形状清单
        inspect = self.client.post("/api/template/inspect", json={"template": payload_id})
        self.assertEqual(200, inspect.status_code)
        shapes = inspect.json()["slides"][0]["shapes"]
        self.assertTrue(any(s["shape_name"] == "COVER_PART_NUMBER" for s in shapes))

        # 5) 显式绑定生成
        gen = self.client.post("/api/template/generate", json={
            "template": payload_id,
            "output_mode": "deck",
            "slides": [{
                "source": 1,
                "bindings": {"COVER_PART_NUMBER": {"type": "text", "source": "f.partNo"}},
            }],
            "data": self.payload,
        })
        self.assertEqual(200, gen.status_code, gen.text)
        self.assertEqual("1", gen.headers["x-dfm-bindings-applied"])
        prs = Presentation(io.BytesIO(gen.content))
        self.assertEqual("TP-HPDC-2026-0087", prs.slides[0].shapes[0].text.strip())

        # 6) 删除
        deleted = self.client.delete(f"/api/templates/{payload_id}")
        self.assertEqual(200, deleted.status_code)

    def test_upload_rejects_non_pptx(self):
        response = self.client.post(
            "/api/templates/upload?template_id=badfile",
            files={"file": ("notes.txt", b"hello", "text/plain")},
        )
        self.assertEqual(415, response.status_code)

    def test_upload_rejects_invalid_zip(self):
        response = self.client.post(
            "/api/templates/upload?template_id=broken",
            files={"file": ("broken.pptx", b"not a zip at all", "application/octet-stream")},
        )
        self.assertEqual(422, response.status_code)

    def test_workbench_page_served(self):
        response = self.client.get("/template-editor")
        self.assertEqual(200, response.status_code)
        self.assertIn("DFM 模板可视化绑定", response.text)
        self.assertIn("btnGenerate", response.text)

    def test_table_demo_scan_recognises_cell_placeholders(self):
        scan = self.client.post("/api/template/scan", json={"template": "table-demo"})
        self.assertEqual(200, scan.status_code)
        self.assertIn("f.partNo", scan.json()["placeholder_paths"])


if __name__ == "__main__":
    unittest.main()
