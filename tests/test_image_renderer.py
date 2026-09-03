# -*- coding: utf-8 -*-
import base64
import io
import unittest

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches

from app.report.ppt.exceptions import RendererError
from app.report.ppt.renderers import ImageRenderMode, ImageRenderer, RenderContext


def png_bytes(width, height):
    stream = io.BytesIO()
    Image.new("RGB", (width, height), "blue").save(stream, format="PNG")
    return stream.getvalue()


def make_frame(width=4, height=4):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(1),
        Inches(1),
        Inches(width),
        Inches(height),
    )
    shape.name = "IMAGE_MAIN"
    return prs, slide, shape


class ImageRendererTests(unittest.TestCase):
    def test_contain_preserves_ratio_centers_and_stays_inside_frame(self):
        prs, slide, frame = make_frame(4, 4)
        bounds = (frame.left, frame.top, frame.width, frame.height)

        result = ImageRenderer().render(
            frame,
            png_bytes(400, 200),
            RenderContext(slide_key="ISSUE_STANDARD", options={"mode": "contain"}),
        )

        self.assertEqual(ImageRenderMode.CONTAIN, result.mode)
        self.assertAlmostEqual(2.0, result.width / result.height, places=3)
        self.assertGreaterEqual(result.left, bounds[0])
        self.assertGreaterEqual(result.top, bounds[1])
        self.assertLessEqual(result.left + result.width, bounds[0] + bounds[2])
        self.assertLessEqual(result.top + result.height, bounds[1] + bounds[3])
        self.assertEqual(1, len(slide.shapes))
        self.assertEqual("IMAGE_MAIN", slide.shapes[0].name)

        output = io.BytesIO()
        prs.save(output)
        output.seek(0)
        reopened = Presentation(output).slides[0].shapes[0]
        self.assertEqual("IMAGE_MAIN", reopened.name)
        self.assertAlmostEqual(2.0, reopened.width / reopened.height, places=3)

    def test_cover_fills_frame_and_center_crops_wide_image(self):
        _, _, frame = make_frame(2, 4)
        bounds = (frame.left, frame.top, frame.width, frame.height)
        result = ImageRenderer().render(
            frame,
            png_bytes(400, 100),
            RenderContext(options={"mode": "cover"}),
        )
        self.assertEqual(bounds, (result.left, result.top, result.width, result.height))
        self.assertGreater(result.crop_left, 0)
        self.assertAlmostEqual(result.crop_left, result.crop_right, places=5)
        self.assertEqual(0, result.crop_top)

    def test_center_crop_crops_tall_image_vertically(self):
        _, _, frame = make_frame(4, 2)
        result = ImageRenderer().render(
            frame,
            png_bytes(100, 400),
            RenderContext(options={"mode": "center_crop"}),
        )
        self.assertEqual(ImageRenderMode.CENTER_CROP, result.mode)
        self.assertGreater(result.crop_top, 0)
        self.assertAlmostEqual(result.crop_top, result.crop_bottom, places=5)

    def test_preserve_aspect_ratio_alias_behaves_as_contain(self):
        _, _, frame = make_frame(3, 3)
        result = ImageRenderer().render(
            frame,
            png_bytes(300, 100),
            RenderContext(options={"mode": "preserve_aspect_ratio"}),
        )
        self.assertEqual(ImageRenderMode.PRESERVE_ASPECT_RATIO, result.mode)
        self.assertAlmostEqual(3.0, result.width / result.height, places=3)

    def test_accepts_base64_data_uri_and_image_list(self):
        _, _, frame = make_frame()
        data_uri = "data:image/png;base64," + base64.b64encode(png_bytes(20, 10)).decode("ascii")
        result = ImageRenderer().render(frame, ["", data_uri], RenderContext())
        self.assertEqual((20, 10), (result.image_width_px, result.image_height_px))

    def test_invalid_mode_does_not_remove_template_shape(self):
        _, slide, frame = make_frame()
        with self.assertRaisesRegex(RendererError, "unsupported image mode=stretch"):
            ImageRenderer().render(frame, png_bytes(10, 10), RenderContext(options={"mode": "stretch"}))
        self.assertEqual(1, len(slide.shapes))
        self.assertEqual("IMAGE_MAIN", slide.shapes[0].name)

    def test_invalid_image_has_contextual_error(self):
        _, _, frame = make_frame()
        with self.assertRaisesRegex(RendererError, "renderer=image; slide=ISSUE_STANDARD; shape=IMAGE_MAIN"):
            ImageRenderer().render(frame, b"not an image", RenderContext(slide_key="ISSUE_STANDARD"))


if __name__ == "__main__":
    unittest.main()
