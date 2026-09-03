# -*- coding: utf-8 -*-
import io
import tempfile
import unittest
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches

from app.dfm.models import SlideType
from app.report.ppt import TemplateLoader, TemplateValidator
from app.report.ppt.exceptions import TemplateValidationError
from app.report.ppt.schema import SlideSchema
from app.report.ppt.schema_loader import SlideSchemaRegistry


def model(data):
    if hasattr(SlideSchema, "model_validate"):
        return SlideSchema.model_validate(data)
    return SlideSchema.parse_obj(data)


def schema_registry(fields, version="1", template="ISSUE_STANDARD"):
    schema = model({
        "schema_version": 1,
        "template_version": version,
        "slide_type": "issue_standard",
        "template": template,
        "fields": fields,
    })
    return SlideSchemaRegistry({SlideType.ISSUE_STANDARD: schema}, version)


def image_stream():
    stream = io.BytesIO()
    Image.new("RGB", (10, 10), "red").save(stream, format="PNG")
    stream.seek(0)
    return stream


def make_template(directory, shapes, version="1", extra_slide=False):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    marker = slide.shapes.add_textbox(Inches(0), Inches(0), Inches(1), Inches(0.2))
    marker.name = "TEMPLATE_KEY__ISSUE_STANDARD"
    for name, kind in shapes:
        if kind == "text":
            shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(0.5))
        elif kind == "rect":
            shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(2), Inches(2), Inches(1))
        elif kind == "picture":
            shape = slide.shapes.add_picture(image_stream(), Inches(1), Inches(3))
        elif kind == "table":
            shape = slide.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(2), Inches(1))
        else:
            raise AssertionError(kind)
        shape.name = name
    if extra_slide:
        extra = prs.slides.add_slide(prs.slide_layouts[6])
        marker = extra.shapes.add_textbox(Inches(0), Inches(0), Inches(1), Inches(0.2))
        marker.name = "TEMPLATE_KEY__UNUSED"
    path = Path(directory) / "template.pptx"
    prs.save(path)
    return TemplateLoader(path, version=version).load()


class TemplateValidatorTests(unittest.TestCase):
    def test_valid_template_passes(self):
        fields = {
            "title": {"source": "title", "shape": "ISSUE_TITLE", "renderer": "text"},
            "image": {"source": "images.main", "shape": "IMAGE_MAIN", "renderer": "image"},
            "table": {"source": "parameters", "shape": "ISSUE_TABLE", "renderer": "table"},
        }
        with tempfile.TemporaryDirectory() as directory:
            template = make_template(directory, [
                ("ISSUE_TITLE", "text"),
                ("IMAGE_MAIN", "rect"),
                ("ISSUE_TABLE", "table"),
            ])
            result = TemplateValidator().validate(template, schema_registry(fields))
        self.assertTrue(result.is_valid)
        self.assertEqual([], result.errors)

    def test_missing_shape_is_an_error_with_binding_context(self):
        fields = {"title": {"source": "title", "shape": "ISSUE_TITLE", "renderer": "text"}}
        with tempfile.TemporaryDirectory() as directory:
            template = make_template(directory, [])
            result = TemplateValidator().validate(template, schema_registry(fields))
        error = result.errors[0]
        self.assertEqual("TEMPLATE_SHAPE_MISSING", error.code)
        self.assertEqual("ISSUE_STANDARD", error.template_key)
        self.assertEqual("ISSUE_TITLE", error.shape)
        self.assertEqual("title", error.field)

    def test_duplicate_shape_is_an_error(self):
        fields = {"title": {"source": "title", "shape": "ISSUE_TITLE", "renderer": "text"}}
        with tempfile.TemporaryDirectory() as directory:
            template = make_template(directory, [("ISSUE_TITLE", "text"), ("ISSUE_TITLE", "text")])
            result = TemplateValidator().validate(template, schema_registry(fields))
        self.assertEqual("TEMPLATE_SHAPE_DUPLICATE", result.errors[0].code)

    def test_incompatible_renderer_shape_is_an_error(self):
        fields = {"title": {"source": "title", "shape": "ISSUE_TITLE", "renderer": "text"}}
        with tempfile.TemporaryDirectory() as directory:
            template = make_template(directory, [("ISSUE_TITLE", "picture")])
            result = TemplateValidator().validate(template, schema_registry(fields))
        self.assertEqual("TEMPLATE_SHAPE_INCOMPATIBLE", result.errors[0].code)
        self.assertEqual("text", result.errors[0].renderer)

    def test_template_version_mismatch_is_an_error(self):
        with tempfile.TemporaryDirectory() as directory:
            template = make_template(directory, [], version="2")
            result = TemplateValidator().validate(template, schema_registry({}, version="1"))
        self.assertIn("TEMPLATE_VERSION_MISMATCH", [item.code for item in result.errors])

    def test_missing_template_page_is_an_error(self):
        with tempfile.TemporaryDirectory() as directory:
            template = make_template(directory, [])
            registry = schema_registry(
                {"title": {"source": "title", "shape": "TITLE", "renderer": "text"}},
                template="SUMMARY",
            )
            result = TemplateValidator().validate(template, registry)
        self.assertEqual("TEMPLATE_SLIDE_MISSING", result.errors[0].code)

    def test_unused_template_page_is_a_warning(self):
        with tempfile.TemporaryDirectory() as directory:
            template = make_template(directory, [], extra_slide=True)
            result = TemplateValidator().validate(template, schema_registry({}))
        self.assertTrue(result.is_valid)
        self.assertEqual("UNUSED_TEMPLATE_SLIDE", result.warnings[0].code)

    def test_validate_or_raise_carries_structured_result(self):
        fields = {"title": {"source": "title", "shape": "MISSING", "renderer": "text"}}
        with tempfile.TemporaryDirectory() as directory:
            template = make_template(directory, [])
            with self.assertRaises(TemplateValidationError) as caught:
                TemplateValidator().validate_or_raise(template, schema_registry(fields))
        self.assertEqual("TEMPLATE_SHAPE_MISSING", caught.exception.result.errors[0].code)


if __name__ == "__main__":
    unittest.main()
