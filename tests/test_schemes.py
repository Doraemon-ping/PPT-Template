# -*- coding: utf-8 -*-
import io
import tempfile
import unittest
from pathlib import Path

from pptx import Presentation

from app.demo import demo_state
from app.report.ppt.scheme_service import SchemeError, SchemeService

ROOT = Path(__file__).resolve().parents[1]

SLIDES = [
    {
        "source": 1,
        "bindings": {
            "b1": {"type": "text", "source": "f.partNo", "shape": "COVER_PART_NUMBER"},
        },
    },
    {"source": 3, "repeat": "t.issues"},
]


class SchemeServiceTests(unittest.TestCase):
    def test_save_list_get_overwrite_delete(self):
        with tempfile.TemporaryDirectory() as tmp:
            service = SchemeService(Path(tmp))
            service.save(name="我的方案", template="demo", slides=SLIDES, description="测试")
            listed = service.list()
            self.assertEqual(1, len(listed))
            self.assertEqual("我的方案", listed[0]["name"])
            self.assertEqual("demo", listed[0]["template"])
            self.assertEqual(2, listed[0]["slide_count"])
            self.assertEqual(1, listed[0]["binding_count"])

            record = service.get("我的方案")
            self.assertEqual(2, len(record["deck"]["slides"]))

            # 同名覆盖 = 更新
            service.save(name="我的方案", template="demo",
                         slides=[{"source": 1}], description="改版")
            self.assertEqual(1, service.list()[0]["slide_count"])

            service.delete("我的方案")
            with self.assertRaises(SchemeError):
                service.get("我的方案")

    def test_save_validations(self):
        with tempfile.TemporaryDirectory() as tmp:
            service = SchemeService(Path(tmp))
            with self.assertRaises(SchemeError):
                service.save(name="  ", template="demo", slides=SLIDES)
            with self.assertRaises(SchemeError):
                service.save(name="方案", template="", slides=SLIDES)
            with self.assertRaises(SchemeError):
                service.save(name="方案", template="demo", slides=[])
            with self.assertRaises(SchemeError):
                service.save(name="方案", template="demo", slides=[{"source": 0}])

    def test_persists_across_instances(self):
        with tempfile.TemporaryDirectory() as tmp:
            SchemeService(Path(tmp)).save(name="跨实例", template="demo", slides=SLIDES)
            reloaded = SchemeService(Path(tmp)).list()
            self.assertEqual("跨实例", reloaded[0]["name"])


try:
    from fastapi.testclient import TestClient
    from app.main import app
except ModuleNotFoundError:  # Local bundled test runtime may omit web dependencies.
    TestClient = None
    app = None


@unittest.skipUnless(TestClient is not None, "FastAPI/httpx test dependencies are not installed")
class SchemeAPITests(unittest.TestCase):
    NAME = "apitest-方案-01"

    @classmethod
    def setUpClass(cls):
        cls.client = TestClient(app)
        cls.payload = demo_state()

    def test_full_flow_save_list_generate_load_delete(self):
        # 保存方案
        save = self.client.post("/api/schemes", json={
            "name": self.NAME,
            "template": "demo",
            "description": "API 测试",
            "slides": [
                {"source": 1, "bindings": {
                    "b1": {"type": "text", "source": "f.partNo", "shape": "COVER_PART_NUMBER"}}},
                {"source": 3, "repeat": "t.issues"},
            ],
        })
        self.assertEqual(200, save.status_code, save.text)
        self.addCleanup(self.client.delete, f"/api/schemes/{self.NAME}")

        # 列表
        listed = self.client.get("/api/schemes").json()["schemes"]
        self.assertTrue(any(item["name"] == self.NAME for item in listed))

        # 载入（完整记录，供二次编辑）
        record = self.client.get(f"/api/schemes/{self.NAME}").json()
        self.assertEqual("demo", record["template"])
        self.assertEqual(2, len(record["deck"]["slides"]))

        # 一键生成：方案 + 当前表单数据
        gen = self.client.post(f"/api/schemes/{self.NAME}/generate", json={"data": self.payload})
        self.assertEqual(200, gen.status_code, gen.text)
        self.assertEqual("openxml-scheme-v1", gen.headers["x-dfm-engine"])
        self.assertEqual("3", gen.headers["x-dfm-slide-count"])  # 封面 + 2 条问题
        prs = Presentation(io.BytesIO(gen.content))
        self.assertEqual(3, len(prs.slides))
        self.assertEqual("TP-HPDC-2026-0087", prs.slides[0].shapes[0].text.strip())

        # 删除
        deleted = self.client.delete(f"/api/schemes/{self.NAME}")
        self.assertEqual(200, deleted.status_code)

    def test_generate_unknown_scheme_404(self):
        response = self.client.post("/api/schemes/not-exist/generate", json={"data": {}})
        self.assertEqual(404, response.status_code)


if __name__ == "__main__":
    unittest.main()
