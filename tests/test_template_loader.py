# -*- coding: utf-8 -*-
import tempfile
import unittest
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

from app.report.ppt.exceptions import (
    TemplateLoadError,
    TemplateNotFoundError,
    TemplateSlideDuplicateError,
    TemplateSlideMissingError,
)
from app.report.ppt.template_loader import TemplateLoader


def add_named_textbox(slide, name):
    shape = slide.shapes.add_textbox(Inches(0), Inches(0), Inches(1), Inches(0.3))
    shape.name = name
    return shape


class TemplateLoaderTests(unittest.TestCase):
    def _save_template(self, directory, keys):
        prs = Presentation()
        for key in keys:
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            add_named_textbox(slide, f"TEMPLATE_KEY__{key}")
        path = Path(directory) / "template.pptx"
        prs.save(path)
        return path

    def test_loads_and_resolves_pages_by_marker_name_not_page_number(self):
        with tempfile.TemporaryDirectory() as directory:
            path = self._save_template(directory, ["SUMMARY", "COVER"])

            loaded = TemplateLoader(path, version="A12", required_keys=["cover", "summary"]).load()

            self.assertEqual("A12", loaded.version)
            self.assertEqual(("SUMMARY", "COVER"), loaded.template_keys)
            self.assertIs(loaded.presentation.slides[1], loaded.get_slide("cover"))

    def test_missing_required_page_fails_during_load(self):
        with tempfile.TemporaryDirectory() as directory:
            path = self._save_template(directory, ["COVER"])
            with self.assertRaisesRegex(TemplateSlideMissingError, "template=SUMMARY"):
                TemplateLoader(path, required_keys=["COVER", "SUMMARY"]).load()

    def test_duplicate_template_key_fails_during_load(self):
        with tempfile.TemporaryDirectory() as directory:
            path = self._save_template(directory, ["COVER", "COVER"])
            with self.assertRaisesRegex(TemplateSlideDuplicateError, "template=COVER"):
                TemplateLoader(path).load()

    def test_missing_file_has_explicit_error(self):
        with tempfile.TemporaryDirectory() as directory:
            with self.assertRaises(TemplateNotFoundError):
                TemplateLoader(Path(directory) / "missing.pptx").load()

    def test_invalid_pptx_has_explicit_error(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "invalid.pptx"
            path.write_text("not a pptx", encoding="utf-8")
            with self.assertRaises(TemplateLoadError):
                TemplateLoader(path).load()


if __name__ == "__main__":
    unittest.main()
