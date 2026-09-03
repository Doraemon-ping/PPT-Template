# -*- coding: utf-8 -*-
import io
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace

from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from app.report.ppt import SlideFactory, TemplateLoader


def image_stream():
    stream = io.BytesIO()
    Image.new("RGB", (40, 20), "purple").save(stream, format="PNG")
    stream.seek(0)
    return stream


class SlideFactoryTests(unittest.TestCase):
    def test_clones_page_repeatedly_preserving_picture_relationships(self):
        with tempfile.TemporaryDirectory() as directory:
            prs = Presentation()
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            marker = slide.shapes.add_textbox(Inches(0), Inches(0), Inches(1), Inches(0.2))
            marker.name = "TEMPLATE_KEY__ISSUE_STANDARD"
            picture = slide.shapes.add_picture(image_stream(), Inches(1), Inches(1), width=Inches(2))
            picture.name = "LOGO"
            path = Path(directory) / "template.pptx"
            prs.save(path)

            loaded = TemplateLoader(path, required_keys=["ISSUE_STANDARD"]).load()
            plans = [
                SimpleNamespace(template_key="ISSUE_STANDARD"),
                SimpleNamespace(template_key="ISSUE_STANDARD"),
            ]
            created = SlideFactory().create(loaded, plans)

            self.assertEqual(2, len(created.presentation.slides))
            self.assertEqual(2, len(created.slides))
            self.assertFalse(any(
                shape.name.startswith("TEMPLATE_KEY__")
                for output_slide in created.slides
                for shape in output_slide.shapes
            ))

            output = io.BytesIO()
            created.presentation.save(output)
            output.seek(0)
            reopened = Presentation(output)
            self.assertEqual(2, len(reopened.slides))
            for output_slide in reopened.slides:
                logo = next(shape for shape in output_slide.shapes if shape.name == "LOGO")
                self.assertGreater(len(logo.image.blob), 0)

    def test_layout_placeholders_are_not_duplicated(self):
        with tempfile.TemporaryDirectory() as directory:
            prs = Presentation()
            slide = prs.slides.add_slide(prs.slide_layouts[0])
            marker = slide.shapes.add_textbox(Inches(0), Inches(0), Inches(1), Inches(0.2))
            marker.name = "TEMPLATE_KEY__COVER"
            source_shape_count_without_marker = len(slide.shapes) - 1
            path = Path(directory) / "template.pptx"
            prs.save(path)

            loaded = TemplateLoader(path, required_keys=["COVER"]).load()
            created = SlideFactory().create(
                loaded,
                [SimpleNamespace(template_key="COVER")],
            )

            self.assertEqual(source_shape_count_without_marker, len(created.slides[0].shapes))


if __name__ == "__main__":
    unittest.main()
