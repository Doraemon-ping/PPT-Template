# -*- coding: utf-8 -*-
import base64
import io
import tempfile
import unittest
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

from app.dfm.models import (
    DFMIssue,
    DFMProject,
    DFMReport,
    DFMReportMetadata,
    DFMSummary,
    IssueSeverity,
)
from app.report.ppt import PPTEngine, SlideSchemaLoader
from app.report.ppt.schema import RendererType


SCHEMA_DIR = Path(__file__).resolve().parents[1] / "app" / "report" / "ppt" / "schemas"


def data_uri(width=200, height=100):
    stream = io.BytesIO()
    Image.new("RGB", (width, height), "orange").save(stream, format="PNG")
    return "data:image/png;base64," + base64.b64encode(stream.getvalue()).decode("ascii")


def add_styled_text(slide, name, x, y):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(2.5), Inches(0.45))
    shape.name = name
    run = shape.text_frame.paragraphs[0].add_run()
    run.text = "TEMPLATE"
    run.font.name = "微软雅黑"
    run.font.size = Pt(14)
    return shape


def build_named_template(path):
    schemas = SlideSchemaLoader().load_directory(SCHEMA_DIR)
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(5.625)
    for schema in schemas.schemas.values():
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        marker = slide.shapes.add_textbox(Inches(0), Inches(0), Inches(0.5), Inches(0.2))
        marker.name = f"TEMPLATE_KEY__{schema.template}"
        for index, binding in enumerate(schema.fields.values()):
            column = index % 3
            row = index // 3
            x = 0.5 + column * 3.1
            y = 0.6 + row * 2.1
            if binding.renderer in {RendererType.IMAGE, RendererType.COMPLEX_IMAGE}:
                shape = slide.shapes.add_shape(
                    MSO_SHAPE.RECTANGLE,
                    Inches(x),
                    Inches(y),
                    Inches(2.8),
                    Inches(1.8),
                )
                shape.name = binding.shape
            else:
                add_styled_text(slide, binding.shape, x, y)
    prs.save(path)


def make_report():
    issue = DFMIssue(
        id="DFM-001",
        category="DraftAngle",
        severity=IssueSeverity.CRITICAL,
        title="拔模角不足",
        description="当前区域拔模角不足",
        recommendation="建议调整至1.5°以上",
        images={"main": data_uri()},
    )
    return DFMReport(
        project=DFMProject(
            project_id="PROJECT-1",
            project_name="测试项目",
            part_name="支架",
            part_number="P-001",
            report_date="2026-09-02",
        ),
        summary=DFMSummary(total=1, critical=1),
        issues=[issue],
        metadata=DFMReportMetadata(report_id="REPORT-1", template_version="1"),
    )


class PPTEngineTests(unittest.TestCase):
    def test_generates_valid_cover_summary_issue_and_conclusion_from_named_template(self):
        with tempfile.TemporaryDirectory() as directory:
            template_path = Path(directory) / "DFM_Master_v1.pptx"
            build_named_template(template_path)

            result = PPTEngine(
                template_path,
                SCHEMA_DIR,
                template_version="1",
            ).generate(make_report())

            self.assertTrue(result.validation.is_valid)
            self.assertEqual(4, len(result.plans))
            self.assertEqual("0.2-migration1", result.generator_version)
            self.assertGreater(len(result.buffer.getvalue()), 0)

            prs = Presentation(result.buffer)
            self.assertEqual(4, len(prs.slides))
            self.assertIn("generator_version=0.2-migration1", prs.core_properties.keywords)
            self.assertIn("report_id=REPORT-1", prs.core_properties.comments)
            self.assertFalse(any(
                shape.name.startswith("TEMPLATE_KEY__")
                for slide in prs.slides
                for shape in slide.shapes
            ))

            issue_slide = prs.slides[2]
            values = {shape.name: shape.text for shape in issue_slide.shapes if shape.has_text_frame}
            self.assertEqual("DFM-001", values["ISSUE_ID"])
            self.assertEqual("拔模角不足", values["ISSUE_TITLE"])
            self.assertEqual("Critical", values["ISSUE_LEVEL"])
            image = next(shape for shape in issue_slide.shapes if shape.name == "IMAGE_MAIN")
            self.assertTrue(image.shape_type)
            self.assertAlmostEqual(2.0, image.width / image.height, places=2)

    def test_engine_can_generate_multiple_pages_from_same_issue_template(self):
        with tempfile.TemporaryDirectory() as directory:
            template_path = Path(directory) / "DFM_Master_v1.pptx"
            build_named_template(template_path)
            report = make_report()
            if hasattr(report.issues[0], "model_copy"):
                second = report.issues[0].model_copy(deep=True)
            else:  # Pydantic v1 compatibility
                second = report.issues[0].copy(deep=True)
            second.id = "DFM-002"
            second.title = "第二个问题"
            report.issues.append(second)
            report.summary.total = 2
            report.summary.critical = 2

            result = PPTEngine(template_path, SCHEMA_DIR).generate(report)
            prs = Presentation(result.buffer)

            self.assertEqual(5, len(prs.slides))
            self.assertEqual("DFM-001", next(s for s in prs.slides[2].shapes if s.name == "ISSUE_ID").text)
            self.assertEqual("DFM-002", next(s for s in prs.slides[3].shapes if s.name == "ISSUE_ID").text)


if __name__ == "__main__":
    unittest.main()
