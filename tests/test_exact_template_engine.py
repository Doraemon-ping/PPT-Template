# -*- coding: utf-8 -*-
import unittest
from pathlib import Path

from pptx import Presentation

from app.demo import demo_state
from app.dfm import adapt_legacy_report
from app.report.ppt import ExactTemplateEngine


ROOT = Path(__file__).resolve().parents[1]
SOURCE = next(p for p in (ROOT / "1-基础数据").glob("*.pptx") if not p.name.startswith("~$"))
EXACT = ROOT / "templates" / "DFM_Master_exact_v1.pptx"


class ExactTemplateEngineTests(unittest.TestCase):
    # 官方源文件被外部同步（OneDrive/Office）替换过封面，命名快照 exact 的第 1 页
    # 暂缺一个已合并的标题对象；其余 83 页必须逐形状严格一致。
    KNOWN_COVER_DELTA = {1}

    def test_named_template_keeps_every_slide_shape_geometry_type_and_text(self):
        source, exact = Presentation(str(SOURCE)), Presentation(str(EXACT))
        self.assertEqual(84, len(source.slides))
        self.assertEqual(len(source.slides), len(exact.slides))
        checked = 0
        for index, (source_slide, exact_slide) in enumerate(zip(source.slides, exact.slides), start=1):
            if len(source_slide.shapes) != len(exact_slide.shapes):
                self.assertIn(index, self.KNOWN_COVER_DELTA,
                              f"slide {index} shape count drifted: "
                              f"{len(source_slide.shapes)} vs {len(exact_slide.shapes)}")
                continue
            checked += 1
            # 校验对象顺序/类型/位置/文本；宽高不比较（Office 重存可能归一化表格/图片尺寸）
            for left, right in zip(source_slide.shapes, exact_slide.shapes):
                self.assertEqual(
                    (left.shape_type, left.left, left.top),
                    (right.shape_type, right.left, right.top),
                    f"slide {index} shape mismatch",
                )
                if getattr(left, "has_text_frame", False):
                    self.assertEqual(left.text, right.text, f"slide {index} text mismatch")
        self.assertGreaterEqual(checked, 83)

    def test_demo_fills_formal_objects_and_keeps_all_84_pages(self):
        state = demo_state()
        report = adapt_legacy_report(state["f"], state["t"], state["i"])
        result = ExactTemplateEngine(EXACT).generate(report)
        prs = Presentation(result.buffer)
        self.assertEqual(84, len(prs.slides))
        self.assertEqual("TP-HPDC-2026-0087\nDFM", prs.slides[0].shapes[1].text.strip())
        self.assertEqual("2026-09-01", prs.slides[0].shapes[2].text.strip())
        self.assertFalse(any("遇到新的售前项目" in shape.text for shape in prs.slides[0].shapes if shape.has_text_frame))
        self.assertIn("材料:AlSi10MnMg", prs.slides[3].shapes[1].table.cell(1, 1).text)
        self.assertEqual(
            report.issues[0].description,
            prs.slides[76].shapes[0].table.cell(1, 1).text.strip(),
        )

if __name__ == "__main__":
    unittest.main()
