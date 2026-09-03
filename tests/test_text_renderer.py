# -*- coding: utf-8 -*-
import io
import unittest

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

from app.report.ppt.exceptions import RendererError
from app.report.ppt.renderers import RenderContext, TextRenderer


def make_text_shape():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    shape.name = "ISSUE_TITLE"
    run = shape.text_frame.paragraphs[0].add_run()
    run.text = "模板标题"
    run.font.name = "微软雅黑"
    run.font.size = Pt(24)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x12, 0x34, 0x56)
    return prs, shape


class TextRendererTests(unittest.TestCase):
    def test_replaces_text_without_changing_template_run_style(self):
        prs, shape = make_text_shape()
        before = shape.text_frame.paragraphs[0].runs[0]
        style = (before.font.name, before.font.size, before.font.bold, before.font.color.rgb)

        result = TextRenderer().render(
            shape,
            "拔模角不足",
            RenderContext(slide_key="ISSUE_STANDARD", field_name="title"),
        )

        after = shape.text_frame.paragraphs[0].runs[0]
        self.assertEqual("拔模角不足", shape.text)
        self.assertEqual(style, (after.font.name, after.font.size, after.font.bold, after.font.color.rgb))
        self.assertFalse(result.truncated)
        self.assertTrue(shape.text_frame.word_wrap)

        output = io.BytesIO()
        prs.save(output)
        output.seek(0)
        reopened = Presentation(output).slides[0].shapes[0].text_frame.paragraphs[0].runs[0]
        self.assertEqual(Pt(24), reopened.font.size)
        self.assertEqual("微软雅黑", reopened.font.name)

    def test_truncates_with_ellipsis_and_reports_possible_overflow(self):
        _, shape = make_text_shape()
        result = TextRenderer().render(
            shape,
            "1234567890",
            RenderContext(options={"max_length": 6}),
        )
        self.assertEqual("12345…", shape.text)
        self.assertTrue(result.truncated)
        self.assertTrue(result.possible_overflow)

    def test_estimates_wrapped_line_overflow_without_shrinking_font(self):
        _, shape = make_text_shape()
        original_size = shape.text_frame.paragraphs[0].runs[0].font.size
        result = TextRenderer().render(
            shape,
            "1234567890",
            RenderContext(options={"chars_per_line": 4, "max_lines": 2}),
        )
        self.assertEqual(3, result.estimated_lines)
        self.assertTrue(result.possible_overflow)
        self.assertEqual(original_size, shape.text_frame.paragraphs[0].runs[0].font.size)

    def test_non_text_shape_has_contextual_error(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        shape = slide.shapes.add_picture(_png_stream(10, 10), Inches(0), Inches(0))
        shape.name = "IMAGE_MAIN"
        with self.assertRaisesRegex(RendererError, "renderer=text; slide=ISSUE_STANDARD; shape=IMAGE_MAIN"):
            TextRenderer().render(shape, "text", RenderContext(slide_key="ISSUE_STANDARD"))


def _png_stream(width, height):
    from PIL import Image

    stream = io.BytesIO()
    Image.new("RGB", (width, height), "red").save(stream, format="PNG")
    stream.seek(0)
    return stream


if __name__ == "__main__":
    unittest.main()
