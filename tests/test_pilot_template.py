# -*- coding: utf-8 -*-
import unittest
from pathlib import Path

from pptx import Presentation

from app.dfm.models import DFMIssue, DFMProject, DFMReport, DFMSummary
from app.demo import demo_state
from app.dfm import adapt_legacy_report
from app.report.ppt import PPTEngine, SlideSchemaLoader, TemplateLoader, TemplateValidator


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "templates" / "DFM_Master_v1.pptx"
SCHEMAS = ROOT / "app" / "report" / "ppt" / "schemas"


class PilotTemplateTests(unittest.TestCase):
    def test_demo_generates_named_part_specifications_page(self):
        state = demo_state()
        report = adapt_legacy_report(state["f"], state["t"], state["i"])

        generated = PPTEngine(TEMPLATE, SCHEMAS, template_version="1").generate(report)
        reopened = Presentation(generated.buffer)

        self.assertEqual(6, len(reopened.slides))
        part_slide = reopened.slides[2]
        values = {
            shape.name: shape.text
            for shape in part_slide.shapes
            if getattr(shape, "has_text_frame", False)
        }
        self.assertEqual("3.85", values["PART_FINISHED_WEIGHT"])
        self.assertEqual("3.0", values["PART_WALL"])
        self.assertEqual("AlSi10MnMg", values["PART_MATERIAL"])
        self.assertEqual("120000", values["PART_ANNUAL_VOLUME"])

    def test_versioned_template_matches_all_schemas(self):
        loaded = TemplateLoader(TEMPLATE, version="1").load()
        schemas = SlideSchemaLoader().load_directory(SCHEMAS)
        result = TemplateValidator().validate(loaded, schemas)

        self.assertTrue(result.is_valid)
        self.assertEqual(
            {
                "COVER", "SUMMARY", "PART_SPECIFICATIONS",
                "ISSUE_STANDARD", "ISSUE_COMPARE", "CONCLUSION",
            },
            set(loaded.template_keys),
        )

    def test_real_pilot_template_generates_reopenable_pptx(self):
        report = DFMReport(
            project=DFMProject(
                project_id="P-1",
                project_name="Pilot",
                part_name="Bracket",
                part_number="PART-001",
                report_date="2026-09-02",
            ),
            summary=DFMSummary(total=1, major=1),
            issues=[DFMIssue(
                id="DFM-001",
                severity="Major",
                title="Pilot issue",
                description="Description",
                recommendation="Recommendation",
            )],
        )

        generated = PPTEngine(TEMPLATE, SCHEMAS, template_version="1").generate(report)
        reopened = Presentation(generated.buffer)

        self.assertEqual(4, len(reopened.slides))
        self.assertTrue(generated.validation.is_valid)
        self.assertEqual(
            ["OPTIONAL_IMAGE_NOT_RENDERED"],
            [warning.code for warning in generated.validation.warnings],
        )


if __name__ == "__main__":
    unittest.main()
