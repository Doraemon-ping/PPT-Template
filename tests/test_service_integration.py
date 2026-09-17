"""Contract, persistence boundaries and real PPTX generation for both forms."""
import io
import json
import sqlite3
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from fastapi.testclient import TestClient
from pptx import Presentation
from pptx.util import Inches
from app.main import app
from app.services import hpdc, machining, workbench, workbench_core
from app.native_forms import key

class ServiceIntegrationTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        for module in (hpdc, machining, workbench, workbench_core):
            p = patch.object(module, 'DATA_DIR', self.root)
            p.start(); self.addCleanup(p.stop)
        self.client = TestClient(app)

    def create_hpdc(self):
        result = self.client.post('/api/form-apps/dfm/projects', json={
            'name': '压铸测试', 'data': {'f': {'partNo': 'HPDC-001', 'aPart': 820, 'castP': 80}, 't': {}, 'i': {}}})
        self.assertEqual(200, result.status_code, result.text)
        return result.json()

    def snapshot(self, source, project):
        result = self.client.get(f'/api/ppt/sources/{source}/projects/{project}/snapshot')
        self.assertEqual(200, result.status_code, result.text)
        return result.json()

    def template_bytes(self, field):
        ppt = Presentation()
        slide = ppt.slides.add_slide(ppt.slide_layouts[6])
        shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1))
        shape.name = 'ReportTitle'; shape.text = '{' + field + '}'
        stream = io.BytesIO(); ppt.save(stream)
        return stream.getvalue()

    def generate(self, source, snapshot, field, expected):
        upload = self.client.post('/api/templates/upload?app_id='+source+'&template_id=contract',
            files={'file': ('contract.pptx', self.template_bytes(field), 'application/vnd.openxmlformats-officedocument.presentationml.presentation')})
        self.assertEqual(200, upload.status_code, upload.text)
        payload = {'name': '契约测试', 'template': 'contract', 'slides': [{'source': 1, 'bindings': {'title': {'shape': 'ReportTitle', 'source': field, 'type': 'text'}}}]}
        self.assertEqual(200, self.client.post('/api/schemes?app_id='+source, json=payload).status_code)
        response = self.client.post('/api/schemes/契约测试/generate?app_id='+source, json={'data': snapshot['data']})
        self.assertEqual(200, response.status_code, response.text[:300] if response.status_code != 200 else '')
        deck = Presentation(io.BytesIO(response.content))
        self.assertEqual(expected, deck.slides[0].shapes[0].text)
        self.assertTrue((self.root/'ppt_workbench'/'sources'/source/'templates'/'contract'/'master.pptx').exists())
        return payload

    def test_hpdc_export_calculation_and_ppt(self):
        project = self.create_hpdc()
        snap = self.snapshot('dfm', project['id'])
        self.assertEqual('1.0', snap['contract_version'])
        self.assertEqual('HPDC-001', snap['data']['f']['partNo'])
        self.assertIn('force', snap['data']['ppt'])
        self.generate('dfm', snap, 'f.partNo', 'HPDC-001')
        self.assertEqual([], self.client.get('/api/templates?app_id=machining-dfm').json()['templates'])

    def test_machining_export_project_changes_and_ppt(self):
        project = self.client.get('/api/machining-dfm/bootstrap').json()['project']
        project['state']['G']['part'] = 'MACH-001'
        updated = self.client.put('/api/machining-dfm/projects/'+project['id'], json={'name': project['name'], 'state': project['state'], 'revision': project['revision']}).json()
        snap = self.snapshot('machining-dfm', project['id'])
        field = 'f.' + key('G.part')
        self.assertEqual(updated['revision'], snap['revision'])
        self.assertNotIn('runtime', snap['data'])
        self.generate('machining-dfm', snap, field, 'MACH-001')
        with machining.store().connect() as db:
            state = json.loads(db.execute('SELECT state_json FROM projects').fetchone()[0])
            self.assertNotIn('mdb', state)
        self.client.post('/api/machining-dfm/projects/'+project['id']+'/archive?archived=true', json={})
        self.assertEqual(410, self.client.get(f"/api/ppt/sources/machining-dfm/projects/{project['id']}/snapshot").status_code)

    def test_drafts_owned_by_workbench_and_conflicts(self):
        self.create_hpdc()
        url='/api/ppt/sources/dfm/drafts/test'
        self.assertEqual(0, self.client.get(url).json()['revision'])
        body={'revision':0,'payload':{'deck':[{'source':1}]}}
        self.assertEqual(200,self.client.put(url,json=body).status_code)
        self.assertEqual(409,self.client.put(url,json=body).status_code)
        db=sqlite3.connect(self.root/'ppt_workbench'/'workbench.sqlite3')
        self.assertEqual(1,db.execute('SELECT count(*) FROM drafts').fetchone()[0]); db.close()
        with hpdc._platform().connect() as db:
            self.assertEqual(0,db.execute('SELECT count(*) FROM drafts').fetchone()[0])

    def test_legacy_raw_generation_calls_form_normalizer(self):
        response=self.client.post('/api/template/formula-preview',json={'expression':'F={ppt.force.part}', 'data':{'f':{'aPart':820,'castP':80}}})
        self.assertEqual(200,response.status_code,response.text[:300] if response.status_code!=200 else '')
        self.assertEqual('0',response.headers['x-dfm-formula-missing'])

    def test_service_boundaries(self):
        ppt=TestClient(workbench.app)
        self.assertEqual(404,ppt.get('/api/machining-dfm/bootstrap').status_code)
        self.assertEqual(404,ppt.post('/api/calc',json={}).status_code)
        self.assertEqual(404,TestClient(machining.app).post('/api/calc',json={}).status_code)
        self.assertEqual(200,TestClient(hpdc.app).post('/api/calc',json={}).status_code)
        self.assertEqual(404,TestClient(hpdc.app).get('/api/machining-dfm/bootstrap').status_code)
        self.assertEqual(404,ppt.get('/api/ppt/sources/unknown').status_code)

    def test_migrate_template_and_scheme_preserves_original(self):
        old=self.root/'machining_dfm'/'data'/'schemes'
        old.mkdir(parents=True); (old/'old.json').write_text('{"name":"old"}',encoding='utf-8')
        new=workbench_core.storage_root('machining-dfm')
        self.assertTrue((new/'schemes'/'old.json').exists())
        self.assertTrue((old/'old.json').exists())
        (new/'schemes'/'old.json').write_text('{"name":"updated"}',encoding='utf-8')
        workbench_core.storage_root('machining-dfm')
        self.assertEqual('updated',json.loads((new/'schemes'/'old.json').read_text())['name'])

    def test_machining_projection_calculates_without_mutating_project(self):
        from app.machining_projection import report_runtime
        state={'G':{'hpd':8,'dpm':20,'avl':.8},'mdb':[{'rapid':30,'tc':2}],
               'pr':[{'nm':'OP10','mi':0,'mc':2,'tl':[{'n':1000,'vf':100,'ln':10,'ps':2,'cn':1,'d':10}]}]}
        result=report_runtime(state)
        self.assertNotIn('_ct',state['pr'][0]['tl'][0])
        row=result['computed']['processes'][0]
        self.assertEqual(12,row['cut_seconds'])
        self.assertEqual(20,row['noncut_seconds'])
        self.assertEqual(32,row['cycle_seconds'])
        self.assertEqual(28800,row['monthly_capacity'])

    def test_provider_authentication(self):
        import os
        with patch.dict(os.environ, {'PPT_PROVIDER_TOKEN':'test-only-token'}):
            client=TestClient(machining.app)
            self.assertEqual(401,client.get('/api/ppt-provider/v1/sources').status_code)
            self.assertEqual(200,client.get('/api/ppt-provider/v1/sources',headers={'Authorization':'Bearer test-only-token'}).status_code)

    def test_workbench_invalid_context_is_validation_error(self):
        result=self.client.post('/api/template/formula-preview',json={'expression':'test','data':None})
        self.assertEqual(422,result.status_code)

if __name__=='__main__':
    unittest.main()
