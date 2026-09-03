# -*- coding: utf-8 -*-
import io
import unittest
from types import SimpleNamespace

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

from app.dfm.models import SlidePlan, SlideType
from app.report.ppt.exceptions import ReportValidationError
from app.report.ppt.schema import SlideSchema
from app.report.ppt.schema_loader import SlideSchemaRegistry
from app.report.ppt.validators import ReportValidator, SlideValidator


def schema_model(fields, slide_type="issue_standard", template="ISSUE_STANDARD"):
    data = {
        "schema_version": 1,
        "template_version": "1",
        "slide_type": slide_type,
        "template": template,
        "fields": fields,
    }
    if hasattr(SlideSchema, "model_validate"):
        return SlideSchema.model_validate(data)
    return SlideSchema.parse_obj(data)


def issue_schema(image_required=True, title_options=None):
    return schema_model({
        "id": {"source": "id", "shape": "ISSUE_ID", "renderer": "text"},
        "title": {
            "source": "title",
            "shape": "ISSUE_TITLE",
            "renderer": "text",
            "options": title_options or {},
        },
        "image": {
            "source": "images.main",
            "shape": "IMAGE_MAIN",
            "renderer": "image",
            "required": image_required,
        },
    })


def registry(schema):
    return SlideSchemaRegistry({schema.slide_type: schema}, "1")


def png_stream(width=200, height=100):
    stream = io.BytesIO()
    Image.new("RGB", (width, height), "green").save(stream, format="PNG")
    stream.seek(0)
    return stream


def add_text(slide, name, text, x=1, y=1, w=3, h=0.5, size=14):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    shape.name = name
    run = shape.text_frame.paragraphs[0].add_run()
    run.text = text
    run.font.size = Pt(size)
    return shape


def valid_presentation():
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(5.625)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_text(slide, "ISSUE_ID", "DFM-001", y=0.5)
    add_text(slide, "ISSUE_TITLE", "拔模角不足", y=1.2)
    picture = slide.shapes.add_picture(png_stream(), Inches(1), Inches(2), width=Inches(4))
    picture.name = "IMAGE_MAIN"
    return prs


def issue_plan(issue_id="DFM-001"):
    return SlidePlan(
        order=1,
        slide_type=SlideType.ISSUE_STANDARD,
        template_key="ISSUE_STANDARD",
        data={},
        issue_id=issue_id,
    )


class OutputValidatorTests(unittest.TestCase):
    def test_valid_generated_slide_passes(self):
        prs = valid_presentation()
        result = ReportValidator().validate(
            prs,
            plans=[issue_plan()],
            schemas=registry(issue_schema()),
        )
        self.assertTrue(result.is_valid)
        self.assertEqual([], result.warnings)

    def test_shape_outside_slide_is_error(self):
        prs = valid_presentation()
        add_text(prs.slides[0], "OUTSIDE", "x", x=9.5, w=1)
        result = SlideValidator(prs.slide_width, prs.slide_height).validate(
            prs.slides[0], slide_index=1
        )
        self.assertIn("SHAPE_OUT_OF_BOUNDS", [item.code for item in result.errors])

    def test_empty_required_title_reports_both_required_and_title_errors(self):
        prs = valid_presentation()
        title = next(shape for shape in prs.slides[0].shapes if shape.name == "ISSUE_TITLE")
        title.text = ""
        result = SlideValidator(prs.slide_width, prs.slide_height).validate(
            prs.slides[0],
            slide_index=1,
            plan=issue_plan(),
            schema=issue_schema(),
        )
        codes = [item.code for item in result.errors]
        self.assertIn("REQUIRED_FIELD_EMPTY", codes)
        self.assertIn("SLIDE_TITLE_EMPTY", codes)

    def test_placeholder_text_and_small_font_are_errors(self):
        prs = valid_presentation()
        title = next(shape for shape in prs.slides[0].shapes if shape.name == "ISSUE_TITLE")
        title.text_frame.paragraphs[0].runs[0].text = "待填写"
        title.text_frame.paragraphs[0].runs[0].font.size = Pt(8)
        result = SlideValidator(prs.slide_width, prs.slide_height).validate(
            prs.slides[0], slide_index=1, plan=issue_plan(), schema=issue_schema()
        )
        codes = [item.code for item in result.errors]
        self.assertIn("TEMPLATE_PLACEHOLDER_REMAINS", codes)
        self.assertIn("TEXT_FONT_BELOW_MINIMUM", codes)

    def test_text_overflow_is_warning(self):
        prs = valid_presentation()
        title = next(shape for shape in prs.slides[0].shapes if shape.name == "ISSUE_TITLE")
        title.text_frame.paragraphs[0].runs[0].text = "1234567890"
        schema = issue_schema(title_options={"chars_per_line": 4, "max_lines": 2})
        result = SlideValidator(prs.slide_width, prs.slide_height).validate(
            prs.slides[0], slide_index=1, plan=issue_plan(), schema=schema
        )
        self.assertIn("TEXT_MAY_OVERFLOW", [item.code for item in result.warnings])

    def test_missing_required_image_is_error(self):
        prs = valid_presentation()
        slide = prs.slides[0]
        picture = next(shape for shape in slide.shapes if shape.name == "IMAGE_MAIN")
        picture._element.getparent().remove(picture._element)
        frame = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(2), Inches(4), Inches(2))
        frame.name = "IMAGE_MAIN"
        result = SlideValidator(prs.slide_width, prs.slide_height).validate(
            slide, slide_index=1, plan=issue_plan(), schema=issue_schema()
        )
        self.assertIn("REQUIRED_FIELD_EMPTY", [item.code for item in result.errors])

    def test_distorted_image_ratio_is_warning(self):
        prs = valid_presentation()
        slide = prs.slides[0]
        picture = next(shape for shape in slide.shapes if shape.name == "IMAGE_MAIN")
        picture.width = Inches(4)
        picture.height = Inches(4)
        result = SlideValidator(prs.slide_width, prs.slide_height).validate(
            slide, slide_index=1, plan=issue_plan(), schema=issue_schema()
        )
        self.assertIn("IMAGE_ASPECT_RATIO_ABNORMAL", [item.code for item in result.warnings])

    def test_issue_without_id_is_error(self):
        prs = valid_presentation()
        result = SlideValidator(prs.slide_width, prs.slide_height).validate(
            prs.slides[0], slide_index=1, plan=issue_plan(issue_id=""), schema=issue_schema()
        )
        self.assertIn("ISSUE_ID_MISSING", [item.code for item in result.errors])

    def test_report_detects_slide_count_and_unknown_slide_type(self):
        prs = valid_presentation()
        unknown = SimpleNamespace(slide_type="mystery", template_key="MYSTERY", issue_id=None)
        result = ReportValidator().validate(prs, plans=[unknown, unknown])
        codes = [item.code for item in result.errors]
        self.assertIn("SLIDE_COUNT_MISMATCH", codes)
        self.assertIn("SLIDE_TYPE_UNRECOGNIZED", codes)

    def test_validate_or_raise_carries_result(self):
        prs = valid_presentation()
        with self.assertRaises(ReportValidationError) as caught:
            ReportValidator().validate_or_raise(prs, plans=[])
        self.assertEqual("SLIDE_COUNT_MISMATCH", caught.exception.result.errors[0].code)


if __name__ == "__main__":
    unittest.main()
