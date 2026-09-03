# -*- coding: utf-8 -*-
"""Build a versioned named-shape pilot template from the formal A12 deck.

This is a one-time template migration utility, not runtime business/rendering
code. It never overwrites the source deck.
"""

import argparse
from pathlib import Path
from types import SimpleNamespace

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from app.report.ppt.slide_factory import SlideFactory
from app.report.ppt.template_loader import LoadedTemplate


BLUE = RGBColor(0x00, 0x70, 0xC0)
DARK_BLUE = RGBColor(0x00, 0x2B, 0x67)
LIGHT_BLUE = RGBColor(0xE9, 0xF4, 0xFB)
LIGHT_GREY = RGBColor(0xF2, 0xF2, 0xF2)
ORANGE = RGBColor(0xFF, 0xC0, 0x00)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BLACK = RGBColor(0x22, 0x22, 0x22)
FONT = "微软雅黑"


def _remove_all_shapes(slide):
    for shape in list(slide.shapes):
        shape._element.getparent().remove(shape._element)


def _add_text(
    slide,
    name,
    text,
    x,
    y,
    w,
    h,
    *,
    size=18,
    bold=False,
    color=DARK_BLUE,
    align=PP_ALIGN.LEFT,
    fill=None,
    line=None,
    margin=6,
):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    shape.name = name
    if fill is None:
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(1.2)
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    frame.margin_left = frame.margin_right = Pt(margin)
    frame.margin_top = frame.margin_bottom = Pt(2)
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return shape


def _add_rect(slide, name, x, y, w, h, *, fill=WHITE, line=BLUE):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(x),
        Inches(y),
        Inches(w),
        Inches(h),
    )
    shape.name = name
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Pt(1.2)
    return shape


def _add_marker(slide, key):
    marker = _add_text(slide, f"TEMPLATE_KEY__{key}", "", 0, 0, 0.05, 0.05, size=12)
    marker.fill.background()
    marker.line.fill.background()


def _build_cover(slide):
    for shape in list(slide.shapes):
        if shape.shape_type != 13:  # keep formal background and logo pictures
            shape._element.getparent().remove(shape._element)
    _add_marker(slide, "COVER")
    _add_text(slide, "DFM_TITLE", "PART-NO", 4.0, 2.55, 5.35, 0.70,
              size=32, bold=True, align=PP_ALIGN.CENTER)
    _add_text(slide, "STATIC_DFM", "DFM", 4.0, 3.20, 5.35, 0.55,
              size=28, bold=True, align=PP_ALIGN.CENTER)
    _add_text(slide, "PROJECT_NAME", "项目名称", 3.65, 3.92, 6.05, 0.40,
              size=16, bold=True, align=PP_ALIGN.CENTER)
    _add_text(slide, "PART_NAME", "零件名称", 3.65, 4.32, 6.05, 0.36,
              size=14, align=PP_ALIGN.CENTER)
    _add_text(slide, "PART_NUMBER", "零件号", 3.65, 4.68, 6.05, 0.34,
              size=12, align=PP_ALIGN.CENTER)
    _add_text(slide, "REPORT_DATE", "2026-01-01", 5.05, 5.72, 3.25, 0.50,
              size=20, bold=True, align=PP_ALIGN.CENTER)


def _build_summary(slide):
    _remove_all_shapes(slide)
    _add_marker(slide, "SUMMARY")
    _add_text(slide, "STATIC_SUMMARY_TITLE", "DFM 问题汇总", 0.55, 0.45, 6.0, 0.60,
              size=26, bold=True)
    cards = [
        ("SUMMARY_TOTAL", "问题总数", 0.65, LIGHT_BLUE, DARK_BLUE),
        ("SUMMARY_CRITICAL", "Critical", 3.85, RGBColor(0xFD, 0xE9, 0xE7), RGBColor(0xC0, 0x00, 0x00)),
        ("SUMMARY_MAJOR", "Major", 7.05, RGBColor(0xFF, 0xF2, 0xCC), RGBColor(0xBF, 0x76, 0x00)),
        ("SUMMARY_MINOR", "Minor", 10.25, RGBColor(0xE2, 0xF0, 0xD9), RGBColor(0x54, 0x82, 0x35)),
    ]
    for name, label, x, fill, color in cards:
        _add_rect(slide, f"CARD_{name}", x, 2.05, 2.45, 2.75, fill=fill, line=color)
        _add_text(slide, f"LABEL_{name}", label, x + 0.15, 2.35, 2.15, 0.48,
                  size=16, bold=True, color=color, align=PP_ALIGN.CENTER)
        _add_text(slide, name, "0", x + 0.15, 3.05, 2.15, 1.10,
                  size=42, bold=True, color=color, align=PP_ALIGN.CENTER)
    _add_text(slide, "STATIC_SUMMARY_NOTE", "按问题严重度自动统计", 0.65, 5.45, 12.05, 0.45,
              size=13, color=RGBColor(0x66, 0x66, 0x66), align=PP_ALIGN.CENTER)


def _build_part_specifications(slide):
    _remove_all_shapes(slide)
    _add_marker(slide, "PART_SPECIFICATIONS")
    _add_text(slide, "STATIC_PART_SPEC_TITLE", "产品基本参数", 0.55, 0.35, 8.0, 0.58,
              size=25, bold=True)
    entries = [
        ("PART_FINISHED_WEIGHT", "成品重量 (kg)"),
        ("PART_CASTING_WEIGHT", "毛坯重量 (kg)"),
        ("PART_RUNNER_WEIGHT", "流道重量 (kg)"),
        ("PART_OVERFLOW_WEIGHT", "渣包重量 (kg)"),
        ("PART_WALL", "基本壁厚 (mm)"),
        ("PART_WALL_MAX", "最大壁厚 (mm)"),
        ("PART_LENGTH", "长度 L (mm)"),
        ("PART_WIDTH", "宽度 W (mm)"),
        ("PART_HEIGHT", "高度 H (mm)"),
        ("PART_CAVITY_COUNT", "模具腔数"),
        ("PART_CASTING_PRESSURE", "铸造压力 (MPa)"),
        ("PART_MATERIAL", "材料牌号"),
        ("PART_ANNUAL_VOLUME", "年产量 (件/年)"),
        ("PART_PROJECT_TYPE", "项目类型"),
        ("PART_SURFACE_REQUIREMENT", "表面要求"),
        ("PART_LEAK_REQUIREMENT", "气密要求"),
    ]
    for index, (name, label) in enumerate(entries):
        column = index % 4
        row = index // 4
        x = 0.55 + column * 3.15
        y = 1.25 + row * 1.30
        _add_text(slide, f"LABEL_{name}", label, x, y, 2.85, 0.38,
                  size=12, bold=True, color=WHITE, fill=DARK_BLUE)
        _add_text(slide, name, "—", x, y + 0.38, 2.85, 0.65,
                  size=15, color=BLACK, align=PP_ALIGN.CENTER, fill=WHITE, line=BLUE)


def _build_issue_standard(slide):
    _remove_all_shapes(slide)
    _add_marker(slide, "ISSUE_STANDARD")
    _add_text(slide, "ISSUE_TITLE", "问题标题", 0.55, 0.30, 9.2, 0.58,
              size=24, bold=True)
    _add_text(slide, "ISSUE_ID", "DFM-000", 10.15, 0.34, 1.35, 0.42,
              size=13, bold=True, align=PP_ALIGN.CENTER, fill=LIGHT_BLUE, line=BLUE)
    _add_text(slide, "ISSUE_LEVEL", "Critical", 11.65, 0.34, 1.10, 0.42,
              size=12, bold=True, color=DARK_BLUE, align=PP_ALIGN.CENTER, fill=ORANGE, line=DARK_BLUE)
    _add_text(slide, "STATIC_DESC", "问题描述", 0.55, 1.12, 4.15, 0.42,
              size=15, bold=True, color=WHITE, fill=DARK_BLUE)
    _add_text(slide, "ISSUE_DESCRIPTION", "问题描述内容", 0.55, 1.54, 4.15, 2.05,
              size=14, color=BLACK, fill=WHITE, line=BLUE)
    _add_text(slide, "STATIC_REC", "优化建议", 0.55, 3.85, 4.15, 0.42,
              size=15, bold=True, color=WHITE, fill=DARK_BLUE)
    _add_text(slide, "ISSUE_RECOMMENDATION", "优化建议内容", 0.55, 4.27, 4.15, 2.15,
              size=14, color=BLACK, fill=WHITE, line=BLUE)
    _add_text(slide, "STATIC_IMAGE", "问题区域 / CAD 截图", 4.95, 1.12, 7.80, 0.42,
              size=15, bold=True, color=WHITE, fill=DARK_BLUE)
    _add_rect(slide, "IMAGE_MAIN", 4.95, 1.54, 7.80, 4.88, fill=LIGHT_GREY, line=BLUE)


def _build_issue_compare(slide):
    _remove_all_shapes(slide)
    _add_marker(slide, "ISSUE_COMPARE")
    _add_text(slide, "ISSUE_TITLE", "问题标题", 0.55, 0.25, 8.8, 0.55,
              size=23, bold=True)
    _add_text(slide, "ISSUE_ID", "DFM-000", 9.75, 0.30, 1.35, 0.40,
              size=13, bold=True, align=PP_ALIGN.CENTER, fill=LIGHT_BLUE, line=BLUE)
    _add_text(slide, "ISSUE_LEVEL", "Critical", 11.35, 0.30, 1.40, 0.40,
              size=12, bold=True, align=PP_ALIGN.CENTER, fill=ORANGE, line=DARK_BLUE)
    _add_text(slide, "STATIC_INFO", "基本信息", 0.55, 1.05, 3.0, 0.42,
              size=15, bold=True, color=WHITE, fill=DARK_BLUE)
    _add_text(slide, "ISSUE_DESCRIPTION", "问题描述", 0.55, 1.47, 3.0, 2.08,
              size=13, color=BLACK, fill=WHITE, line=BLUE)
    _add_text(slide, "ISSUE_RECOMMENDATION", "优化建议", 0.55, 3.82, 3.0, 2.38,
              size=13, color=BLACK, fill=WHITE, line=BLUE)
    _add_text(slide, "STATIC_BEFORE", "优化前", 3.80, 1.05, 4.25, 0.42,
              size=15, bold=True, color=WHITE, fill=DARK_BLUE, align=PP_ALIGN.CENTER)
    _add_rect(slide, "IMAGE_BEFORE", 3.80, 1.47, 4.25, 4.73, fill=LIGHT_GREY, line=BLUE)
    _add_text(slide, "STATIC_AFTER", "优化后", 8.30, 1.05, 4.45, 0.42,
              size=15, bold=True, color=WHITE, fill=DARK_BLUE, align=PP_ALIGN.CENTER)
    _add_rect(slide, "IMAGE_AFTER", 8.30, 1.47, 4.45, 4.73, fill=LIGHT_GREY, line=BLUE)


def _build_conclusion(slide):
    _add_marker(slide, "CONCLUSION")
    entries = [
        ("SUMMARY_TOTAL", "Total", 2.25, DARK_BLUE),
        ("SUMMARY_CRITICAL", "Critical", 5.05, RGBColor(0xC0, 0x00, 0x00)),
        ("SUMMARY_MAJOR", "Major", 7.85, RGBColor(0xBF, 0x76, 0x00)),
        ("SUMMARY_MINOR", "Minor", 10.65, RGBColor(0x54, 0x82, 0x35)),
    ]
    for name, label, x, color in entries:
        _add_text(slide, f"LABEL_{name}", label, x, 6.56, 1.15, 0.28,
                  size=12, bold=True, color=color, align=PP_ALIGN.CENTER)
        _add_text(slide, name, "0", x, 6.88, 1.15, 0.34,
                  size=16, bold=True, color=color, align=PP_ALIGN.CENTER)


def build(source: Path, output: Path):
    source = source.resolve()
    output = output.resolve()
    if source == output:
        raise ValueError("output must not overwrite source template")
    presentation = Presentation(str(source))
    if len(presentation.slides) < 84:
        raise ValueError(f"expected at least 84 slides, got {len(presentation.slides)}")
    selected = LoadedTemplate(
        path=source,
        version="1",
        presentation=presentation,
        slides={
            "COVER": presentation.slides[0],
            "SUMMARY": presentation.slides[75],
            "PART_SPECIFICATIONS": presentation.slides[3],
            "ISSUE_STANDARD": presentation.slides[76],
            "ISSUE_COMPARE": presentation.slides[76],
            "CONCLUSION": presentation.slides[83],
        },
    )
    plans = [
        SimpleNamespace(template_key="COVER"),
        SimpleNamespace(template_key="SUMMARY"),
        SimpleNamespace(template_key="PART_SPECIFICATIONS"),
        SimpleNamespace(template_key="ISSUE_STANDARD"),
        SimpleNamespace(template_key="ISSUE_COMPARE"),
        SimpleNamespace(template_key="CONCLUSION"),
    ]
    created = SlideFactory().create(selected, plans)
    _build_cover(created.slides[0])
    _build_summary(created.slides[1])
    _build_part_specifications(created.slides[2])
    _build_issue_standard(created.slides[3])
    _build_issue_compare(created.slides[4])
    _build_conclusion(created.slides[5])
    created.presentation.core_properties.title = "DFM Named-Shape Master Template"
    created.presentation.core_properties.subject = "template_version=1"
    output.parent.mkdir(parents=True, exist_ok=True)
    created.presentation.save(str(output))


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    build(args.source, args.output)


if __name__ == "__main__":
    main()
