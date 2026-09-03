# -*- coding: utf-8 -*-
"""Validate one generated slide for content and layout defects."""

import math
from typing import Iterable, Optional

from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Pt

from ....dfm.models import SlideType
from ..exceptions import TemplateShapeDuplicateError, TemplateShapeMissingError
from ..schema import RendererType
from ..shape_resolver import ShapeResolver
from .result import ValidationResult


DEFAULT_PLACEHOLDER_PATTERNS = (
    "图片占位",
    "待填写",
    "placeholder",
    "click to add",
    "点击此处",
    "{{",
)


class SlideValidator:
    def __init__(
        self,
        slide_width: int,
        slide_height: int,
        *,
        min_font_size_pt: float = 12.0,
        image_ratio_tolerance: float = 0.05,
        placeholder_patterns: Iterable[str] = DEFAULT_PLACEHOLDER_PATTERNS,
        shape_resolver=None,
    ) -> None:
        self.slide_width = slide_width
        self.slide_height = slide_height
        self.min_font_size_pt = float(min_font_size_pt)
        self.image_ratio_tolerance = float(image_ratio_tolerance)
        self.placeholder_patterns = tuple(pattern.casefold() for pattern in placeholder_patterns)
        self.shape_resolver = shape_resolver or ShapeResolver()

    def validate(
        self,
        slide,
        *,
        slide_index: int,
        plan=None,
        schema=None,
    ) -> ValidationResult:
        result = ValidationResult()
        template_key = getattr(plan, "template_key", None)
        issue_id = getattr(plan, "issue_id", None)
        self._validate_boundaries(result, slide, slide_index, template_key, issue_id)
        self._validate_placeholder_text(result, slide, slide_index, template_key, issue_id)

        if plan is not None and self._is_issue_plan(plan) and not str(issue_id or "").strip():
            result.add_error(
                "ISSUE_ID_MISSING",
                "issue slide has no issue ID",
                slide_index=slide_index,
                template_key=template_key,
            )

        if schema is not None:
            for field_name, binding in schema.fields.items():
                self._validate_binding(
                    result,
                    slide,
                    slide_index=slide_index,
                    template_key=template_key or schema.template,
                    issue_id=issue_id,
                    field_name=field_name,
                    binding=binding,
                )
        return result

    def _validate_boundaries(self, result, slide, slide_index, template_key, issue_id):
        for shape in slide.shapes:
            if shape.name.startswith("TEMPLATE_KEY__"):
                continue
            outside = (
                shape.left < 0
                or shape.top < 0
                or shape.left + shape.width > self.slide_width
                or shape.top + shape.height > self.slide_height
            )
            if outside:
                result.add_error(
                    "SHAPE_OUT_OF_BOUNDS",
                    f"shape exceeds slide boundary: {shape.name}",
                    slide_index=slide_index,
                    template_key=template_key,
                    shape=shape.name,
                    issue_id=issue_id,
                    details={
                        "left": shape.left,
                        "top": shape.top,
                        "width": shape.width,
                        "height": shape.height,
                        "slide_width": self.slide_width,
                        "slide_height": self.slide_height,
                    },
                )

    def _validate_placeholder_text(self, result, slide, slide_index, template_key, issue_id):
        for shape in self.shape_resolver.iter_shapes(slide):
            if not getattr(shape, "has_text_frame", False):
                continue
            text = shape.text.strip()
            folded = text.casefold()
            if text and any(pattern in folded for pattern in self.placeholder_patterns):
                result.add_error(
                    "TEMPLATE_PLACEHOLDER_REMAINS",
                    f"template placeholder text remains in shape: {shape.name}",
                    slide_index=slide_index,
                    template_key=template_key,
                    shape=shape.name,
                    issue_id=issue_id,
                    details={"text": text[:200]},
                )

    def _validate_binding(
        self,
        result,
        slide,
        *,
        slide_index,
        template_key,
        issue_id,
        field_name,
        binding,
    ):
        try:
            shape = self.shape_resolver.get_shape(slide, binding.shape, slide_key=template_key)
        except TemplateShapeMissingError:
            result.add_error(
                "OUTPUT_SHAPE_MISSING",
                f"generated slide is missing bound shape: {binding.shape}",
                slide_index=slide_index,
                template_key=template_key,
                shape=binding.shape,
                field=field_name,
                renderer=binding.renderer.value,
                issue_id=issue_id,
            )
            return
        except TemplateShapeDuplicateError:
            result.add_error(
                "OUTPUT_SHAPE_DUPLICATE",
                f"generated slide contains duplicate bound shape: {binding.shape}",
                slide_index=slide_index,
                template_key=template_key,
                shape=binding.shape,
                field=field_name,
                renderer=binding.renderer.value,
                issue_id=issue_id,
            )
            return

        has_content = self._has_content(shape, binding.renderer)
        if binding.required and not has_content:
            result.add_error(
                "REQUIRED_FIELD_EMPTY",
                f"required field is empty: {field_name}",
                slide_index=slide_index,
                template_key=template_key,
                shape=shape.name,
                field=field_name,
                renderer=binding.renderer.value,
                issue_id=issue_id,
            )
        if field_name.casefold() == "title" and not has_content:
            result.add_error(
                "SLIDE_TITLE_EMPTY",
                "slide title is empty",
                slide_index=slide_index,
                template_key=template_key,
                shape=shape.name,
                field=field_name,
                issue_id=issue_id,
            )

        if binding.renderer in {RendererType.TEXT, RendererType.BADGE} and getattr(shape, "has_text_frame", False):
            self._validate_text(result, shape, binding, slide_index, template_key, issue_id, field_name)
        elif binding.renderer in {RendererType.IMAGE, RendererType.COMPLEX_IMAGE}:
            self._validate_image(result, shape, binding, slide_index, template_key, issue_id, field_name)

    def _validate_text(self, result, shape, binding, slide_index, template_key, issue_id, field_name):
        text = shape.text or ""
        options = dict(binding.options)
        max_length = self._positive_int(options.get("max_length"))
        max_lines = self._positive_int(options.get("max_lines"))
        chars_per_line = self._positive_int(options.get("chars_per_line"))
        estimated_lines = self._estimate_lines(text, chars_per_line)
        if (max_length is not None and len(text) > max_length) or (
            max_lines is not None and estimated_lines > max_lines
        ):
            result.add_warning(
                "TEXT_MAY_OVERFLOW",
                f"text may overflow shape: {shape.name}",
                slide_index=slide_index,
                template_key=template_key,
                shape=shape.name,
                field=field_name,
                renderer=binding.renderer.value,
                issue_id=issue_id,
                details={"length": len(text), "estimated_lines": estimated_lines},
            )
        if max_length is not None and text.endswith("…") and len(text) >= max_length:
            result.add_warning(
                "TEXT_TRUNCATED",
                f"text appears truncated: {shape.name}",
                slide_index=slide_index,
                template_key=template_key,
                shape=shape.name,
                field=field_name,
                renderer=binding.renderer.value,
                issue_id=issue_id,
            )

        minimum = float(options.get("min_font_size", self.min_font_size_pt))
        sizes = [
            run.font.size
            for paragraph in shape.text_frame.paragraphs
            for run in paragraph.runs
            if run.text and run.font.size is not None
        ]
        too_small = [size for size in sizes if size < Pt(minimum)]
        if too_small:
            result.add_error(
                "TEXT_FONT_BELOW_MINIMUM",
                f"text font is below minimum {minimum:g}pt: {shape.name}",
                slide_index=slide_index,
                template_key=template_key,
                shape=shape.name,
                field=field_name,
                renderer=binding.renderer.value,
                issue_id=issue_id,
                details={"minimum_pt": minimum, "actual_pt": min(size.pt for size in too_small)},
            )

    def _validate_image(self, result, shape, binding, slide_index, template_key, issue_id, field_name):
        if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
            if not binding.required:
                result.add_warning(
                    "OPTIONAL_IMAGE_NOT_RENDERED",
                    f"optional image was not rendered: {shape.name}",
                    slide_index=slide_index,
                    template_key=template_key,
                    shape=shape.name,
                    field=field_name,
                    renderer=binding.renderer.value,
                    issue_id=issue_id,
                )
            return
        image_width, image_height = shape.image.size
        visible_width = image_width * max(0.0, 1.0 - shape.crop_left - shape.crop_right)
        visible_height = image_height * max(0.0, 1.0 - shape.crop_top - shape.crop_bottom)
        if visible_width <= 0 or visible_height <= 0 or shape.height <= 0:
            result.add_error(
                "IMAGE_RATIO_INVALID",
                f"image has invalid crop or dimensions: {shape.name}",
                slide_index=slide_index,
                template_key=template_key,
                shape=shape.name,
                field=field_name,
                renderer=binding.renderer.value,
                issue_id=issue_id,
            )
            return
        source_ratio = visible_width / visible_height
        displayed_ratio = shape.width / shape.height
        distortion = abs(displayed_ratio / source_ratio - 1.0)
        if distortion > self.image_ratio_tolerance:
            result.add_warning(
                "IMAGE_ASPECT_RATIO_ABNORMAL",
                f"image display ratio differs from source ratio: {shape.name}",
                slide_index=slide_index,
                template_key=template_key,
                shape=shape.name,
                field=field_name,
                renderer=binding.renderer.value,
                issue_id=issue_id,
                details={
                    "source_ratio": source_ratio,
                    "displayed_ratio": displayed_ratio,
                    "distortion": distortion,
                },
            )

    @staticmethod
    def _has_content(shape, renderer):
        if renderer in {RendererType.TEXT, RendererType.BADGE}:
            return bool(getattr(shape, "has_text_frame", False) and shape.text.strip())
        if renderer in {RendererType.IMAGE, RendererType.COMPLEX_IMAGE}:
            return shape.shape_type == MSO_SHAPE_TYPE.PICTURE
        if renderer == RendererType.TABLE:
            if not getattr(shape, "has_table", False):
                return False
            return any(cell.text.strip() for row in shape.table.rows for cell in row.cells)
        return False

    @staticmethod
    def _is_issue_plan(plan):
        value = getattr(getattr(plan, "slide_type", None), "value", getattr(plan, "slide_type", None))
        return value in {SlideType.ISSUE_STANDARD.value, SlideType.ISSUE_COMPARE.value}

    @staticmethod
    def _positive_int(value):
        try:
            number = int(value)
        except (TypeError, ValueError):
            return None
        return number if number > 0 else None

    @staticmethod
    def _estimate_lines(text, chars_per_line):
        lines = text.splitlines() or [""]
        if chars_per_line is None:
            return len(lines)
        return sum(max(1, math.ceil(len(line) / chars_per_line)) for line in lines)
