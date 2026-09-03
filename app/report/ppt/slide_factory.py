# -*- coding: utf-8 -*-
"""Clone named template pages without depending on page indexes."""

from copy import deepcopy
from dataclasses import dataclass
from typing import List

from pptx.oxml.ns import qn

from .exceptions import SlideFactoryError
from .shape_resolver import ShapeResolver


_SKIPPED_RELATIONSHIP_TYPES = {
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/slideLayout",
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/notesSlide",
}
_RELATIONSHIP_ATTRIBUTES = (qn("r:embed"), qn("r:link"), qn("r:id"))


@dataclass(frozen=True)
class CreatedSlides:
    presentation: object
    slides: List[object]


class SlideFactory:
    """Create output slides by cloning template pages in SlidePlan order."""

    def __init__(self, marker_prefix: str = "TEMPLATE_KEY__") -> None:
        self.marker_prefix = marker_prefix
        self.shape_resolver = ShapeResolver()

    def create(self, template, plans) -> CreatedSlides:
        presentation = template.presentation
        original_slide_ids = list(presentation.slides._sldIdLst)
        created = []
        for plan in plans:
            try:
                source = template.get_slide(plan.template_key)
                destination = self._clone_slide(presentation, source)
                self._remove_template_markers(destination)
                created.append(destination)
            except Exception as exc:
                if isinstance(exc, SlideFactoryError):
                    raise
                raise SlideFactoryError(plan.template_key, str(exc)) from exc
        self._remove_original_slides(presentation, original_slide_ids)
        return CreatedSlides(presentation=presentation, slides=created)

    def _clone_slide(self, presentation, source):
        destination = presentation.slides.add_slide(source.slide_layout)
        # add_slide() materializes layout placeholders. Source slide shapes are
        # copied next, so remove these auto-created placeholders to avoid twins.
        for generated_shape in list(destination.shapes):
            parent = generated_shape._element.getparent()
            if parent is not None:
                parent.remove(generated_shape._element)
        relationship_map = self._copy_relationships(source, destination)
        for shape in source.shapes:
            element = deepcopy(shape.element)
            self._rewrite_relationship_ids(element, relationship_map)
            destination.shapes._spTree.insert_element_before(element, "p:extLst")

        source_c_sld = source._element.cSld
        destination_c_sld = destination._element.cSld
        if source_c_sld.bg is not None:
            if destination_c_sld.bg is not None:
                destination_c_sld.remove(destination_c_sld.bg)
            destination_c_sld.insert(0, deepcopy(source_c_sld.bg))
        if source._element.clrMapOvr is not None:
            if destination._element.clrMapOvr is not None:
                destination._element.remove(destination._element.clrMapOvr)
            destination._element.append(deepcopy(source._element.clrMapOvr))
        return destination

    @staticmethod
    def _copy_relationships(source, destination):
        relationship_map = {}
        for old_id, relationship in source.part.rels.items():
            if relationship.reltype in _SKIPPED_RELATIONSHIP_TYPES:
                continue
            new_id = destination.part.rels._add_relationship(
                relationship.reltype,
                relationship._target,
                relationship.is_external,
            )
            relationship_map[old_id] = new_id
        return relationship_map

    @staticmethod
    def _rewrite_relationship_ids(element, relationship_map):
        for descendant in element.iter():
            for attribute in _RELATIONSHIP_ATTRIBUTES:
                old_id = descendant.get(attribute)
                if old_id in relationship_map:
                    descendant.set(attribute, relationship_map[old_id])

    def _remove_template_markers(self, slide):
        markers = [
            shape
            for shape in self.shape_resolver.iter_shapes(slide)
            if shape.name.startswith(self.marker_prefix)
        ]
        for marker in markers:
            parent = marker._element.getparent()
            if parent is not None:
                parent.remove(marker._element)

    @staticmethod
    def _remove_original_slides(presentation, original_slide_ids):
        slide_id_list = presentation.slides._sldIdLst
        for slide_id in original_slide_ids:
            relationship_id = slide_id.rId
            slide_id_list.remove(slide_id)
            presentation.part.drop_rel(relationship_id)
