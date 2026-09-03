# -*- coding: utf-8 -*-
"""Load a PPTX template and index template pages by named marker shapes.

Each reusable template page is identified by a shape whose PowerPoint name is
``TEMPLATE_KEY__<KEY>``. The marker can be a normal invisible/off-canvas shape;
its text and position are irrelevant. Page numbers are deliberately not part of
the public contract, so template pages may be reordered safely.
"""

from dataclasses import dataclass
from pathlib import Path
from typing import Dict, Iterable, Mapping, Optional

from pptx import Presentation

from .exceptions import (
    TemplateLoadError,
    TemplateNotFoundError,
    TemplateSlideDuplicateError,
    TemplateSlideMissingError,
)
from .shape_resolver import ShapeResolver


DEFAULT_MARKER_PREFIX = "TEMPLATE_KEY__"


@dataclass(frozen=True)
class LoadedTemplate:
    path: Path
    version: str
    presentation: object
    slides: Mapping[str, object]

    def get_slide(self, template_key: str):
        key = _normalize_key(template_key)
        try:
            return self.slides[key]
        except KeyError as exc:
            raise TemplateSlideMissingError(key, self.path) from exc

    @property
    def template_keys(self):
        return tuple(self.slides.keys())


class TemplateLoader:
    """Read-only loader for versioned, marker-indexed PowerPoint templates."""

    def __init__(
        self,
        path,
        *,
        version: str = "1",
        marker_prefix: str = DEFAULT_MARKER_PREFIX,
        required_keys: Optional[Iterable[str]] = None,
    ) -> None:
        self.path = Path(path)
        self.version = str(version)
        self.marker_prefix = marker_prefix
        self.required_keys = tuple(_normalize_key(key) for key in (required_keys or ()))

    def load(self) -> LoadedTemplate:
        if not self.path.is_file():
            raise TemplateNotFoundError(self.path)
        try:
            presentation = Presentation(str(self.path))
        except Exception as exc:  # python-pptx exposes several package/XML errors
            raise TemplateLoadError(self.path, str(exc)) from exc

        slides = self._index_template_slides(presentation)
        for key in self.required_keys:
            if key not in slides:
                raise TemplateSlideMissingError(key, self.path)
        return LoadedTemplate(
            path=self.path.resolve(),
            version=self.version,
            presentation=presentation,
            slides=slides,
        )

    def _index_template_slides(self, presentation) -> Dict[str, object]:
        indexed: Dict[str, object] = {}
        resolver = ShapeResolver()
        for slide in presentation.slides:
            keys = {
                _normalize_key(shape.name[len(self.marker_prefix):])
                for shape in resolver.iter_shapes(slide)
                if shape.name.startswith(self.marker_prefix)
                and shape.name[len(self.marker_prefix):].strip()
            }
            for key in keys:
                if key in indexed:
                    raise TemplateSlideDuplicateError(key, self.path)
                indexed[key] = slide
        return indexed


def _normalize_key(value: str) -> str:
    return str(value).strip().upper()
