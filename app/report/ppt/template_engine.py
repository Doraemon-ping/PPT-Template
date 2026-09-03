# -*- coding: utf-8 -*-
"""Template-driven PPT generation engine (pptx-template style).

Pipeline:
  1. Open the template pptx at the zip level (every untouched part keeps its
     original bytes).
  2. Expand the declarative deck against the data context (``{f,t,i}`` plus
     derived calculations and verdict texts).
  3. For every planned slide: clone the source template slide, fill
     ``{path}`` placeholders and swap bound images.
  4. Rebuild ``presentation.xml`` in deck mode, or fill in place otherwise.

The engine never calls ``python-pptx`` to re-save the official template, never
runs PowerPoint COM, and never recomputes DFM conclusions (they come from
``app.calc`` as data).
"""
import io
import time
from dataclasses import dataclass, field, replace
from pathlib import Path
from typing import Any, Dict, List, Mapping, Optional

from .deck import DeckDefinition, DeckPlanner, SlideOp
from .openxml.image_binding import ImageBindingFiller
from .openxml.package_editor import OoxmlPackage, OoxmlPackageError
from .openxml.placeholder_scanner import PlaceholderScanner
from .openxml.slide_repeater import clone_slide, rebuild_presentation
from .openxml.text_binding import PathResolver, PlaceholderResolutionError, TextBindingFiller
from ...calc import calc_force, compute_all
from ...machines import cur_machine

GENERATOR_VERSION = "1.0-openxml1"


class TemplateEngineError(RuntimeError):
    pass


@dataclass(frozen=True)
class TemplateEngineResult:
    buffer: bytes
    slide_count: int
    template: str
    generator_version: str
    generation_time_seconds: float
    stats: Mapping[str, Any] = field(default_factory=dict)
    warnings: tuple = ()

    @property
    def bytes_io(self) -> io.BytesIO:
        stream = io.BytesIO(self.buffer)
        stream.seek(0)
        return stream


class TemplateEngine:
    """Generate a report pptx from a registered template, a deck, and data."""

    def __init__(
        self,
        template_path,
        *,
        generator_version: str = GENERATOR_VERSION,
    ) -> None:
        if isinstance(template_path, (str, Path)):
            self.template_path = Path(template_path)
        else:  # 也支持内存中的 bytes / BytesIO（测试用）
            self.template_path = template_path
        self.generator_version = generator_version

    # ------------------------------------------------------------------ main
    def generate(
        self,
        deck: DeckDefinition,
        data: Mapping[str, Any],
        *,
        template_map: Optional[Mapping[str, Any]] = None,
        preview_first_item: bool = False,
        preview_page: Optional[int] = None,
    ) -> TemplateEngineResult:
        """Generate with a deck. ``template_map`` maps template ids to paths when
        the deck mixes pages from several templates; the deck's own template is
        the base output package."""
        started = time.perf_counter()
        if deck.output_mode not in {"deck", "in_place"}:
            raise TemplateEngineError(f"unsupported output_mode: {deck.output_mode}")
        paths: Dict[str, Any] = dict(template_map or {})
        if deck.template not in paths:
            paths[deck.template] = self.template_path
        try:
            base_package = OoxmlPackage(paths[deck.template])
        except OoxmlPackageError as exc:
            raise TemplateEngineError(f"template load failed: {exc}") from exc

        context = build_data_context(data)
        plan = DeckPlanner().expand(deck, context)
        if preview_first_item:
            plan = replace(plan, ops=tuple(replace(op, items=op.items[:1]) for op in plan.ops))
        from .openxml.table_pagination import paginate_op
        packages = {deck.template: base_package}
        expanded = []
        try:
            for op in plan.ops:
                template_id = op.template or deck.template
                if template_id not in packages:
                    if template_id not in paths:
                        raise TemplateEngineError(f'模板未提供路径：{template_id}')
                    packages[template_id] = OoxmlPackage(paths[template_id])
                if op.source_index not in packages[template_id].slide_parts():
                    raise TemplateEngineError(f'模板 {template_id} 缺少页面：{op.source_index}')
                expanded.extend(paginate_op(packages[template_id], op, context))
        except (ValueError, RuntimeError) as exc:
            raise TemplateEngineError(f'表格分页失败：{exc}') from exc
        preview_count = len(expanded)
        selected_page = min(max(1, preview_page or 1), max(1, preview_count))
        plan = replace(plan, ops=tuple(expanded[selected_page-1:selected_page] if preview_page is not None else expanded))
        stats: Dict[str, Any] = {
            "slides": 0,
            "text_replaced": 0,
            "text_missing": [],
            "images": 0,
            "images_missing": [],
            "bindings_applied": 0,
            "bindings_skipped": 0,
        }
        filler = TextBindingFiller(missing=deck.missing)

        def _op_template(op) -> str:
            return op.template or deck.template

        if deck.output_mode == "in_place":
            targets = [(op.template or deck.template, op.source_index) for op in plan.ops]
            if len(set(targets)) != len(targets):
                raise TemplateEngineError('内容需要续页，请使用 deck 模式生成（原地填充不能增加页面）')
            for op in plan.ops:
                if _op_template(op) != deck.template:
                    raise TemplateEngineError("in_place 模式仅支持单一模板")
                if len(op.items) != 1 or op.is_repeat:
                    raise TemplateEngineError("in_place mode does not support repeat slides")
                slide_parts = base_package.slide_parts()
                if op.source_index not in slide_parts:
                    raise TemplateEngineError(
                        f"deck references missing template slides: {op.source_index}"
                    )
                part = slide_parts[op.source_index]
                self._fill_part(base_package, part, op, op.items[0], context, filler, stats)
                stats["slides"] += 1
        else:
            # 快速路径（PowerPoint 兼容）：单一模板、无重复页、每源页只用一次时，
            # 原地填充并保留所需原页、删除其余（不克隆/不导入；官方等复杂模板
            # 的“克隆页”会被 PowerPoint 判为损坏）。
            single_template = all(_op_template(op) == deck.template for op in plan.ops)
            no_repeat = all(len(op.items) <= 1 and not op.is_repeat for op in plan.ops)
            source_uses: Dict[int, int] = {}
            for op in plan.ops:
                source_uses[op.source_index] = source_uses.get(op.source_index, 0) + 1
            if single_template and no_repeat and all(v == 1 for v in source_uses.values()):
                slide_parts = base_package.slide_parts()
                kept: List[str] = []
                for op in plan.ops:
                    if op.source_index not in slide_parts:
                        raise TemplateEngineError(
                            f"deck references missing template slides: {op.source_index}"
                        )
                    part = slide_parts[op.source_index]
                    self._fill_part(base_package, part, op, op.items[0], context, filler, stats)
                    kept.append(part)
                    stats["slides"] += 1
                rebuild_presentation(base_package, kept)
            else:
                from .openxml.slide_importer import SlideImporter

                packages: Dict[str, OoxmlPackage] = {deck.template: base_package}
                importer = SlideImporter()
                output_parts: List[str] = []

                for op in plan.ops:
                    template_id = _op_template(op)
                    path = paths.get(template_id)
                    if path is None:
                        raise TemplateEngineError(f"模板未提供路径：{template_id}")
                    if template_id not in packages:
                        try:
                            packages[template_id] = OoxmlPackage(path)
                        except OoxmlPackageError as exc:
                            raise TemplateEngineError(
                                f"template load failed: {template_id} -> {path}"
                            ) from exc
                    package = packages[template_id]
                    slide_parts = package.slide_parts()
                    if op.source_index not in slide_parts:
                        raise TemplateEngineError(
                            f"模板 {template_id} 缺少页面：{op.source_index}"
                        )
                    source_part = slide_parts[op.source_index]
                    for item in op.items:
                        if template_id == deck.template:
                            base = packages[deck.template]
                            new_part = f"ppt/slides/slide{base.next_slide_number()}.xml"
                            clone_slide(package, source_part, new_part)
                        else:
                            new_part = importer.import_slide(
                                packages[deck.template], package, source_part
                            )
                        self._fill_part(
                            packages[deck.template], new_part, op, item, context, filler, stats
                        )
                        output_parts.append(new_part)
                        stats["slides"] += 1
                rebuild_presentation(packages[deck.template], output_parts)

        output = base_package.save()
        slide_count = self._verify_slide_count(output)
        elapsed = time.perf_counter() - started
        return TemplateEngineResult(
            buffer=output,
            slide_count=slide_count,
            template=deck.template,
            generator_version=self.generator_version,
            generation_time_seconds=round(elapsed, 4),
            stats={
                "output_slides": stats["slides"],
                "text_replaced": stats["text_replaced"],
                "text_missing": sorted(set(stats["text_missing"])),
                "images_bound": stats["images"],
                "images_missing": stats["images_missing"],
                "bindings_applied": stats["bindings_applied"],
                "bindings_skipped": stats["bindings_skipped"],
                "missing_policy": deck.missing,
                "preview_page_count": preview_count,
                "preview_page": selected_page,
            },
            warnings=tuple(plan.warnings) + tuple(
                f"图片数据未填写，保留模板原内容：{path}" for path in stats["images_missing"]
            ),
        )

    # ---------------------------------------------------------------- helpers
    def _fill_part(
        self,
        package: OoxmlPackage,
        part: str,
        op: SlideOp,
        item: Optional[Mapping[str, Any]],
        context: Mapping[str, Any],
        filler: TextBindingFiller,
        stats: Dict[str, Any],
    ) -> None:
        scopes = ([item] if item is not None else []) + [context]
        resolver = PathResolver(tuple(scopes))
        xml = package.read(part)
        try:
            new_xml, fill_stats = filler.fill(xml, resolver)
        except PlaceholderResolutionError as exc:
            raise TemplateEngineError(
                f"slide {part} placeholder error: {exc}"
            ) from exc
        stats["text_replaced"] += fill_stats.replaced
        stats["text_missing"].extend(fill_stats.missing)
        package.write(part, new_xml)

        if op.images:
            image_filler = ImageBindingFiller()
            for shape_name, path in op.images.items():
                value, found = resolver.resolve(path)
                if not found or _is_empty(value):
                    continue
                slide_index = _part_index(part)
                image_filler.fill(package, slide_index, shape_name, value)
                stats["images"] += 1

        if op.bindings:
            self._apply_shape_bindings(package, part, op, resolver, stats, original_xml=xml)

    def _apply_shape_bindings(self, package, part, op, resolver, stats, original_xml=None) -> None:
        """Apply explicit bindings, including composite text and visual regions."""
        from lxml import etree

        from .openxml.shape_binding import (
            ShapeBindingError,
            find_shape,
            set_shape_text,
            set_table_cell,
            set_table_rows,
        )
        from .openxml.visual_binding import VisualBindingFiller, render_data_template

        xml = package.read(part)
        root = etree.fromstring(xml)
        errors = []
        changed = False
        visual_tasks = []
        for bind_key, spec in op.bindings.items():
            target_name = spec.shape or bind_key
            shape_id = spec.options.get("shape_id")
            value, found = resolver.resolve(spec.source)
            # Older editors saved image fields on cells as text. Normalize this
            # unambiguous legacy case without changing the user's saved scheme.
            sample = value[0] if isinstance(value, (list, tuple)) and value else value
            if spec.type == 'table_cell' and (spec.source.startswith('i.') or
                    isinstance(sample, str) and sample.startswith('data:image/')):
                spec = spec.model_copy(update={'type': 'image_region'})
            table_slice = op.table_data.get(bind_key)
            if table_slice is not None:
                value, found = table_slice[0], True
            if spec.type in {"image", "image_region"} and (not found or _is_empty(value)):
                # Missing form images are normal during report preparation. Keep
                # the source object intact; corrupt supplied images still fail below.
                stats["bindings_skipped"] += 1
                stats["images_missing"].append(spec.source)
                continue
            if not found:
                if spec.type == "table_rows":
                    # 整表填入：表数据尚未填写时跳过（不报错、不清表），填入后自动生效
                    stats["bindings_skipped"] += 1
                    continue
                if spec.required:
                    errors.append(f"slide {part}: binding source missing: {spec.source}")
                continue
            if spec.type == 'text_replace':
                from copy import deepcopy
                from .openxml.partial_text import replace_text_segments, text_target
                try:
                    shape = find_shape(root, target_name, shape_id=shape_id, context=part)
                    original_root = etree.fromstring(original_xml)
                    original_shape = find_shape(original_root, target_name, shape_id=shape_id, context=part)
                    replacement = deepcopy(text_target(original_shape, spec.options))
                    replace_text_segments(replacement, spec.options, resolver)
                    target = text_target(shape, spec.options)
                    target.getparent().replace(target, replacement)
                    changed = True
                    stats['bindings_applied'] += 1
                except ShapeBindingError as exc:
                    errors.append(f'slide {part}: {exc}')
            elif spec.type in {"text", "text_template", "table_cell", "table_rows"}:
                is_empty = value is None or (isinstance(value, str) and not value.strip())
                if spec.type != "table_rows" and is_empty and spec.options.get("empty") == "keep":
                    continue
                try:
                    shape = find_shape(root, target_name, shape_id=shape_id, context=part)
                    if spec.type == "table_cell":
                        set_table_cell(shape, int(spec.options["row"]), int(spec.options["column"]), value)
                    elif spec.type == "table_rows":
                        records = value if isinstance(value, (list, tuple)) else (
                            [] if value is None else [value])
                        options = dict(spec.options or {})
                        columns_map = options.get("columns_map")
                        set_table_rows(
                            shape,
                            [r for r in records if isinstance(r, dict)],
                            columns_keys=list(options.get("columns_keys") or []),
                            columns_map=dict(columns_map) if columns_map else None,
                            start_row=int(options.get("start_row", 1)),
                            uniform_height=table_slice[1] if table_slice is not None else None,
                            style_row=table_slice[2] if table_slice is not None else None,
                        )
                    else:
                        rendered = value
                        if spec.type == "text_template":
                            rendered = render_data_template(spec.options.get("template", ""), resolver)
                        if "row" in spec.options and "column" in spec.options:
                            set_table_cell(
                                shape, int(spec.options["row"]), int(spec.options["column"]), rendered
                            )
                        else:
                            set_shape_text(shape, "" if rendered is None else str(rendered))
                    changed = True
                    stats["bindings_applied"] += 1
                except ShapeBindingError as exc:
                    if spec.required:
                        errors.append(f"slide {part}: 绑定失败（对象 {target_name}，字段 {spec.source}，类型 {spec.type}）：{exc}")
                    else:
                        stats["bindings_skipped"] += 1
            elif spec.type in {"image", "image_region", "formula"}:
                if _is_empty(value):
                    continue
                visual_tasks.append((target_name, shape_id, spec, value))
            else:
                errors.append(f"slide {part}: unsupported binding type: {spec.type}")
        if errors:
            raise TemplateEngineError("; ".join(errors))
        if changed:
            package.write(part, etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True))
        visual = VisualBindingFiller()
        for target_name, shape_id, spec, value in visual_tasks:
            try:
                if spec.type == "image":
                    ImageBindingFiller().fill(
                        package, _part_index(part), target_name, value, shape_id=shape_id
                    )
                elif spec.type == "image_region":
                    fill = visual.fill_table_region if 'row' in spec.options and 'column' in spec.options else visual.fill_region
                    fill(
                        package, _part_index(part), target_name, value,
                        shape_id=shape_id, options=spec.options,
                    )
                else:
                    formula = render_data_template(spec.options.get("template", ""), resolver)
                    visual.fill_formula(
                        package, _part_index(part), target_name, formula,
                        shape_id=shape_id, options=spec.options,
                    )
                stats["images"] += 1
                stats["bindings_applied"] += 1
            except Exception as exc:  # noqa: BLE001
                label = "图片绑定失败" if spec.type in {"image", "image_region"} else "公式绑定失败"
                message = f"slide {part}: {label}（形状 {target_name}，字段 {spec.source}）：{exc}"
                if spec.required:
                    errors.append(message)
                else:
                    stats["bindings_skipped"] += 1
        if errors:
            raise TemplateEngineError("; ".join(errors))

    @staticmethod
    def _verify_slide_count(output: bytes) -> int:
        try:
            from pptx import Presentation

            return len(Presentation(io.BytesIO(output)).slides)
        except Exception:  # python-pptx is a lenient checker only.
            return 0


def build_data_context(data: Mapping[str, Any]) -> Dict[str, Any]:
    """Merge raw ``{f,t,i}`` with derived values and calc verdict texts."""
    f = dict(data.get("f") or {})
    t = dict(data.get("t") or {})
    i = dict(data.get("i") or {})
    calculated = compute_all(f, apply_machine=False)
    verdicts = {key: value.get("verdict", "") for key, value in calculated["results"].items()}
    force = calc_force(f)
    machine = cur_machine(f)

    def display(value, digits=0):
        try:
            number = float(value)
        except (TypeError, ValueError):
            return "" if value is None else str(value)
        if digits == 0:
            return str(int(round(number)))
        return f"{number:.{digits}f}".rstrip("0").rstrip(".")

    safety = float(f.get("lockSafetyFactor") or 1.1)
    mold_structure = str(f.get("moldStructure") or f"{display(force['cav'])}出{display(force['cav'])}")
    product_info = "\n".join([
        f"零件号:{f.get('partNo', '')};",
        f"成品重量:{f.get('wFinish', '')}kg",
        f"毛坯重量:{f.get('wCast', '')}kg;",
        f"流道重量:{f.get('wRunner', '')}kg",
        f"渣包重量:{f.get('wOverflow', '')}kg",
        f"基本壁厚:{f.get('wall', '')}mm;",
        f"产品尺寸:{f.get('dimL', '')}X{f.get('dimW', '')}X{f.get('dimH', '')}mm;",
        f"高压压铸设备吨位：{display(machine.get('ton'))}T；模具结构{mold_structure}；",
        f"铸造压力:{f.get('castP', '')}MPa;高真空压铸",
        f"材料：{f.get('material', '')}",
    ])
    ppt_force = {
        "pressure": display(force["P"]),
        "area_part": display(force["aPart"]),
        "area_slider": display(force["aSl"]),
        "area_runner": display(force["aRunner"]),
        "area_overflow": display(force["aOver"]),
        "part": display(force["fPart"]),
        "slider": display(force["fSl"]),
        "runner": display(force["fRunner"]),
        "overflow": display(force["fOver"]),
        "total": display(force["fTot"]),
        "slider_angle": display(force["sliderAngle"]),
        "slider_term": (
            f" × tan {display(force['sliderAngle'])}°" if force["sliderAngle"] else ""
        ),
        "clamp_factor": display(safety, 2),
        "clamp_required": display(force["fTot"] * safety),
    }
    return {
        "f": f,
        "t": t,
        "i": i,
        "derived": calculated["derived"],
        "calc_results": verdicts,
        "ppt": {"product_info": product_info, "force": ppt_force},
    }


def _part_index(part: str) -> int:
    import re

    match = re.match(r"ppt/slides/slide(\d+)\.xml", part)
    return int(match.group(1)) if match else 0


def _is_empty(value: Any) -> bool:
    if value is None:
        return True
    if isinstance(value, str):
        return not value.strip()
    if isinstance(value, (list, tuple, dict)):
        return not value
    return False


# Re-exported for callers that only need the high-level API.
scan_placeholders = PlaceholderScanner().scan
