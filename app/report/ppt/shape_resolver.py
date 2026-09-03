# -*- coding: utf-8 -*-
"""Resolve template objects exclusively by PowerPoint Shape Name."""

from typing import Iterator, List

from pptx.enum.shapes import MSO_SHAPE_TYPE

from .exceptions import TemplateShapeDuplicateError, TemplateShapeMissingError


class ShapeResolver:
    """Strict named-shape lookup, including shapes nested in group shapes."""

    def get_shape(self, slide, name: str, *, slide_key: str = "UNKNOWN"):
        matches = self.find_shapes(slide, name)
        if not matches:
            raise TemplateShapeMissingError(slide=slide_key, shape=name)
        if len(matches) > 1:
            raise TemplateShapeDuplicateError(slide=slide_key, shape=name)
        return matches[0]

    def find_shapes(self, slide, name: str) -> List[object]:
        """Return all exact-name matches without silently choosing a duplicate."""

        return [shape for shape in self.iter_shapes(slide) if shape.name == name]

    def has_shape(self, slide, name: str) -> bool:
        return bool(self.find_shapes(slide, name))

    def iter_shapes(self, slide) -> Iterator[object]:
        for shape in slide.shapes:
            yield shape
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                yield from self._iter_group_shapes(shape)

    def _iter_group_shapes(self, group_shape) -> Iterator[object]:
        for shape in group_shape.shapes:
            yield shape
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                yield from self._iter_group_shapes(shape)


def get_shape(slide, name: str, *, slide_key: str = "UNKNOWN"):
    """Convenience API matching ``get_shape(slide, 'ISSUE_TITLE')`` usage."""

    return ShapeResolver().get_shape(slide, name, slide_key=slide_key)
